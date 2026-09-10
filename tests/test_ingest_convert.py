"""Tests for :mod:`src.ingest.convert`.

Every fixture is authored in-test rather than committed as a binary blob: a
checked-in PDF/DOCX is unreviewable in a diff, and the exact bytes matter here
(page order, scrambled content streams, a text layer that is deliberately
absent). ``_build_pdf`` hand-writes a conforming PDF — no authoring library
exists in this repo's dependency set, and pypdfium2 reads PDFs but cannot write
them — and ``_write_docx`` builds a minimal OOXML package with ``zipfile`` when
``python-docx`` is not installed, so the markitdown route is always exercised
rather than skipped.
"""

from __future__ import annotations

import io
import sys
import types
import zipfile
from pathlib import Path

import pytest

from src.ingest.convert import (
    DEFAULT_MAX_CHARS,
    PAGE_BREAK,
    ConversionError,
    ConvertResult,
    MissingConversionDependency,
    UnsupportedConversionFormat,
    convert_to_markdown,
)

pypdfium2 = pytest.importorskip("pypdfium2", reason="extraction extra not installed")
pytest.importorskip("markitdown", reason="extraction extra not installed")


# ----------------------------------------------------------------- fixtures


def _build_pdf(pages: list[list[tuple[str, float, float]]]) -> bytes:
    """Hand-write a PDF with a real text layer.

    ``pages`` is one list of ``(text, x, y)`` placements per page, emitted into
    the content stream in the order given — which is what lets a test control
    whether PDFium's native text order is reading order or scrambled.
    """
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    pages_num = len(objects) + 1
    objects.append(b"")  # placeholder, patched once the kids are known

    page_nums: list[int] = []
    for placements in pages:
        ops = [b"BT /F1 12 Tf"]
        for text, x, y in placements:
            escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
            ops.append(f"1 0 0 1 {x} {y} Tm ({escaped}) Tj".encode("latin-1"))
        ops.append(b"ET")
        stream = b"\n".join(ops)
        content = add(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        page_nums.append(
            add(
                b"<< /Type /Page /Parent "
                + str(pages_num).encode()
                + b" 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 "
                + str(font).encode()
                + b" 0 R >> >> /Contents "
                + str(content).encode()
                + b" 0 R >>"
            )
        )

    kids = b" ".join(str(num).encode() + b" 0 R" for num in page_nums)
    objects[pages_num - 1] = b"<< /Type /Pages /Count " + str(len(page_nums)).encode() + b" /Kids [" + kids + b"] >>"
    catalog = add(b"<< /Type /Catalog /Pages " + str(pages_num).encode() + b" 0 R >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0] * (len(objects) + 1)
    for num, body in enumerate(objects, start=1):
        offsets[num] = len(out)
        out += str(num).encode() + b" 0 obj\n" + body + b"\nendobj\n"
    xref_offset = len(out)
    out += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n0000000000 65535 f \n"
    for num in range(1, len(objects) + 1):
        out += f"{offsets[num]:010d} 00000 n \n".encode()
    out += (
        b"trailer\n<< /Size "
        + str(len(objects) + 1).encode()
        + b" /Root "
        + str(catalog).encode()
        + b" 0 R >>\nstartxref\n"
        + str(xref_offset).encode()
        + b"\n%%EOF\n"
    )
    return bytes(out)


_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

_PACKAGE_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""

_DOCUMENT_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Agnes Handbook</w:t></w:r></w:p>
<w:p><w:r><w:t>Revenue is recognised on delivery.</w:t></w:r></w:p>
</w:body></w:document>"""


def _write_docx(path: Path) -> Path:
    """Write a two-paragraph .docx, with or without ``python-docx`` present."""
    try:
        import docx  # type: ignore[import-not-found]
    except ImportError:
        with zipfile.ZipFile(path, "w") as package:
            package.writestr("[Content_Types].xml", _CONTENT_TYPES)
            package.writestr("_rels/.rels", _PACKAGE_RELS)
            package.writestr("word/document.xml", _DOCUMENT_XML)
        return path

    document = docx.Document()
    document.add_heading("Agnes Handbook", level=1)
    document.add_paragraph("Revenue is recognised on delivery.")
    document.save(str(path))
    return path


#: Same VML (``w:pict``/``v:imagedata``) shape real Word writes for an
#: embedded picture — the one mammoth's ``_read_blip``/``read_imagedata``
#: recognizes with the fewest namespaces to hand-declare. The image bytes
#: are never validated by mammoth (it only opens and base64-encodes them),
#: so a placeholder body is enough; ``python-docx`` is not installed in this
#: environment (see the module docstring), so this is always hand-written.
def _write_docx_with_images(path: Path, sections: list[tuple[str, str]]) -> Path:
    """A ``.docx`` with one ``Heading1`` + one embedded picture per
    ``(heading_text, relationship_id)`` pair in ``sections``, in order."""
    body_parts = []
    rels_parts = []
    for heading, rel_id in sections:
        body_parts.append(
            f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>{heading}</w:t></w:r></w:p>'
            f'<w:p><w:r><w:pict><v:shape><v:imagedata r:id="{rel_id}"/></v:shape></w:pict></w:r></w:p>'
        )
        rels_parts.append(
            f'<Relationship Id="{rel_id}" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            f'Target="media/{rel_id}.png"/>'
        )
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        "<w:document "
        'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        'xmlns:v="urn:schemas-microsoft-com:vml">'
        "<w:body>" + "".join(body_parts) + "</w:body></w:document>"
    )
    rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(rels_parts)
        + "</Relationships>"
    )
    with zipfile.ZipFile(path, "w") as package:
        package.writestr("[Content_Types].xml", _CONTENT_TYPES)
        package.writestr("_rels/.rels", _PACKAGE_RELS)
        package.writestr("word/document.xml", document_xml)
        package.writestr("word/_rels/document.xml.rels", rels_xml)
        for _heading, rel_id in sections:
            package.writestr(f"word/media/{rel_id}.png", b"not a real png; mammoth never validates the bytes")
    return path


def _tiny_png_bytes() -> bytes:
    """A minimal, real 1x1 PNG. Unlike mammoth's docx image handler
    (never inspects the bytes), python-pptx's picture part reads the PNG
    header for native size/aspect ratio, so the pptx fixture below needs
    genuine image bytes rather than a placeholder."""
    import struct
    import zlib

    def _chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00\x00\x00\x00")
    return sig + _chunk(b"IHDR", ihdr) + _chunk(b"IDAT", idat) + _chunk(b"IEND", b"")


def _write_pptx_with_images(path: Path, slide_count: int) -> Path:
    """A real ``.pptx`` (python-pptx — already a hard markitdown dependency,
    never a new one for this test module) with one embedded picture on each
    of ``slide_count`` slides."""
    from pptx import Presentation
    from pptx.util import Inches

    png = _tiny_png_bytes()
    presentation = Presentation()
    layout = presentation.slide_layouts[6]  # blank layout
    for _ in range(slide_count):
        slide = presentation.slides.add_slide(layout)
        slide.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(1))
    presentation.save(str(path))
    return path


# --------------------------------------------------------------- passthrough


@pytest.mark.parametrize(
    "name, mime",
    [
        ("notes.md", "text/markdown"),
        ("notes.txt", "text/plain"),
        ("rows.csv", "text/csv"),
        ("blob.json", "application/json"),
        ("conf.yaml", "text/yaml"),
        ("conf.yml", "text/yaml"),
    ],
)
def test_text_suffixes_pass_through_verbatim(tmp_path, name, mime):
    body = "# Heading\n\nline one\nline two\n"
    path = tmp_path / name
    path.write_text(body, encoding="utf-8")

    result = convert_to_markdown(path, mime)

    assert isinstance(result, ConvertResult)
    assert result.engine == "passthrough"
    assert result.markdown == body


def test_passthrough_replaces_undecodable_bytes_instead_of_failing(tmp_path):
    path = tmp_path / "latin.txt"
    path.write_bytes(b"caf\xe9 revenue")

    result = convert_to_markdown(path, "text/plain")

    assert result.engine == "passthrough"
    assert "revenue" in result.markdown
    assert "�" in result.markdown


def test_passthrough_uses_mime_when_the_name_has_no_suffix(tmp_path):
    path = tmp_path / "README"
    path.write_text("plain body", encoding="utf-8")

    assert convert_to_markdown(path, "text/plain; charset=utf-8").engine == "passthrough"


def test_blank_text_file_is_reported_as_empty(tmp_path):
    path = tmp_path / "blank.txt"
    path.write_text("   \n\n\t\n", encoding="utf-8")

    result = convert_to_markdown(path, "text/plain")

    assert result == ConvertResult(markdown="", engine="empty")


# ---------------------------------------------------------------- size guard


def test_max_chars_truncates_with_a_marker(tmp_path):
    path = tmp_path / "big.txt"
    path.write_text("x" * 500, encoding="utf-8")

    result = convert_to_markdown(path, "text/plain", max_chars=100)

    assert result.engine == "passthrough"
    body, _, marker = result.markdown.partition("\n\n[truncated")
    assert body == "x" * 100
    assert "exceeded 100 characters" in marker


def test_output_at_the_cap_is_not_marked_truncated(tmp_path):
    path = tmp_path / "exact.txt"
    path.write_text("y" * 100, encoding="utf-8")

    result = convert_to_markdown(path, "text/plain", max_chars=100)

    assert result.markdown == "y" * 100


def test_max_chars_must_be_positive(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("body", encoding="utf-8")

    with pytest.raises(ValueError):
        convert_to_markdown(path, "text/plain", max_chars=0)


def test_default_max_chars_is_five_million():
    assert DEFAULT_MAX_CHARS == 5_000_000


# ----------------------------------------------------------------------- pdf


def test_pdf_extracts_text_in_reading_order(tmp_path):
    path = tmp_path / "hello.pdf"
    path.write_bytes(_build_pdf([[("Hello Agnes", 72, 700), ("Second line", 72, 660)]]))

    result = convert_to_markdown(path, "application/pdf")

    assert result.engine == "pypdfium2"
    assert result.markdown == "Hello Agnes\n\nSecond line"


def test_pdf_pages_are_separated_by_a_horizontal_rule(tmp_path):
    path = tmp_path / "two.pdf"
    path.write_bytes(
        _build_pdf(
            [
                [("Page one body", 72, 700)],
                [("Page two body", 72, 700)],
            ]
        )
    )

    result = convert_to_markdown(path, "application/pdf")

    assert result.engine == "pypdfium2"
    assert result.markdown == f"Page one body{PAGE_BREAK}Page two body"


def test_scrambled_content_stream_is_sorted_top_to_bottom_left_to_right(tmp_path):
    # Emitted bottom-first, and right-before-left on the top line: PDFium's
    # native order is "the lower line, then the upper line".
    path = tmp_path / "scrambled.pdf"
    path.write_bytes(
        _build_pdf(
            [
                [
                    ("bottom of the page", 72, 600),
                    ("rightmost column", 300, 700),
                    ("leftmost column", 72, 700),
                ]
            ]
        )
    )

    result = convert_to_markdown(path, "application/pdf")

    assert result.engine == "pypdfium2"
    lines = [line for line in result.markdown.splitlines() if line]
    assert lines[0] == "leftmost column rightmost column"
    assert lines[1] == "bottom of the page"


def test_native_reading_order_is_preserved_not_re_sorted(tmp_path):
    path = tmp_path / "ordered.pdf"
    path.write_bytes(
        _build_pdf(
            [
                [
                    ("alpha heading", 72, 700),
                    ("beta paragraph", 72, 660),
                    ("gamma footer", 72, 620),
                ]
            ]
        )
    )

    result = convert_to_markdown(path, "application/pdf")

    assert result.markdown == "alpha heading\n\nbeta paragraph\n\ngamma footer"


def test_pdf_without_a_text_layer_is_empty_not_an_error(tmp_path):
    path = tmp_path / "scan.pdf"
    path.write_bytes(_build_pdf([[], []]))

    result = convert_to_markdown(path, "application/pdf")

    assert result == ConvertResult(markdown="", engine="empty")


def test_pdf_with_a_near_empty_text_layer_is_empty(tmp_path):
    # A scanned page frequently yields a stray page number and nothing else.
    path = tmp_path / "near_empty.pdf"
    path.write_bytes(_build_pdf([[("7", 300, 60)]]))

    assert convert_to_markdown(path, "application/pdf").engine == "empty"


def test_pdf_uses_mime_when_the_name_has_no_suffix(tmp_path):
    path = tmp_path / "downloaded"
    path.write_bytes(_build_pdf([[("Hello Agnes from Graph", 72, 700)]]))

    result = convert_to_markdown(path, "application/pdf")

    assert result.engine == "pypdfium2"
    assert "Hello Agnes" in result.markdown


def test_pdf_output_is_truncated_at_max_chars(tmp_path):
    page = [("word " * 40, 20, 700 - 20 * row) for row in range(20)]
    path = tmp_path / "long.pdf"
    path.write_bytes(_build_pdf([page, page, page]))

    result = convert_to_markdown(path, "application/pdf", max_chars=200)

    assert result.engine == "pypdfium2"
    assert "truncated" in result.markdown


# --------------------------------------------------------------- markitdown


def test_docx_routes_to_markitdown(tmp_path):
    path = _write_docx(tmp_path / "handbook.docx")

    result = convert_to_markdown(
        path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert result.engine == "markitdown"
    assert "Agnes Handbook" in result.markdown
    assert "Revenue is recognised on delivery." in result.markdown


def test_html_routes_to_markitdown(tmp_path):
    path = tmp_path / "page.html"
    path.write_text(
        "<html><body><h1>Quarterly</h1><p>EMEA grew.</p></body></html>",
        encoding="utf-8",
    )

    result = convert_to_markdown(path, "text/html")

    assert result.engine == "markitdown"
    assert "Quarterly" in result.markdown
    assert "EMEA grew." in result.markdown


def test_html_relative_image_is_not_falsely_disclosed_as_lost(tmp_path):
    """markitdown converts ``<img src="logo.jpg">`` to ``![Logo](logo.jpg)``
    — the IDENTICAL bare-filename shape python-pptx's OWN lost-picture
    placeholder uses (see the module docstring's "Embedded pictures"
    section). Nothing was dropped here: the image reference is exactly what
    the source HTML said, and this document never showed a PowerPoint
    slide-number marker at all. Rewriting it into "not indexed" would be a
    FALSE disclosure — live finding 2026-09-09, the regression this test
    guards."""
    path = tmp_path / "page.html"
    path.write_text(
        '<html><body><h1>Quarterly</h1><img src="logo.jpg" alt="Logo"><p>EMEA grew.</p></body></html>',
        encoding="utf-8",
    )

    result = convert_to_markdown(path, "text/html")

    assert result.engine == "markitdown"
    assert result.image_count == 0
    assert "![Logo](logo.jpg)" in result.markdown
    assert "not indexed" not in result.markdown


# ------------------------------------------------------- legacy office (LibreOffice)


#: Legacy suffix → the OOXML target format LibreOffice must produce, per the
#: mapping in :mod:`src.ingest.convert`.
_LEGACY_OFFICE_CASES = [
    (".doc", "docx"),
    (".rtf", "docx"),
    (".odt", "docx"),
    (".ppt", "pptx"),
    (".odp", "pptx"),
    (".xls", "xlsx"),
    (".ods", "xlsx"),
    (".xlsb", "xlsx"),
    (".xlsm", "xlsx"),
]


def _ooxml_shaped_bytes(note: bytes = b"a package no reader accepts") -> bytes:
    """A real zip package whose CONTENT no reader accepts.

    The OOXML archive guard refuses a ``.docx``/``.pptx``/``.xlsx`` that is
    not a zip AT ALL before any route runs — nothing a LibreOffice resave or a
    CSV/PDF rescue rung does can turn non-zip bytes into an OOXML package, so
    the guard answers for free what three subprocesses would answer slowly. A
    test that wants to exercise what happens AFTER a reader rejects the file
    therefore has to hand it a structurally valid zip. That is also the shape
    of the live finding this chain exists for: a genuine ``.xlsx`` openpyxl
    refuses, not a file that was never a workbook. Same convention as the
    dependency-probe fixture below.
    """
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        package.writestr("[Content_Types].xml", "<Types/>")
        package.writestr("agnes-test-note.txt", note)
    return buffer.getvalue()


def _stub_soffice(
    monkeypatch,
    convert_module,
    *,
    returncode=0,
    produce_output=True,
    side_effect=None,
    output_bytes: bytes = b"fake converted bytes",
):
    """Replace ``soffice`` with a fake that never shells out for real.

    Records every invocation's argv and the temp ``--outdir`` it was given
    (so a test can assert the dir is gone afterwards), and — unless told
    otherwise — drops a placeholder output file at the path LibreOffice
    itself would have written, so the caller's glob for the converted file
    succeeds without a real LibreOffice on the machine. ``output_bytes``
    defaults to an inert placeholder (fine when the caller of the converted
    file is ALSO mocked, e.g. ``_convert_markitdown``) — a rescue-chain test
    reading the converted file for REAL (the CSV/PDF fallback rungs, which
    read openpyxl/pypdfium2 directly rather than through a mock) passes real
    xlsx/PDF bytes here instead.
    """
    calls: list[dict] = []

    def _which(cmd):
        return "/usr/bin/soffice" if cmd == "soffice" else None

    def _run(argv, **kwargs):
        if side_effect is not None:
            raise side_effect
        outdir = Path(argv[argv.index("--outdir") + 1])
        target_format = argv[argv.index("--convert-to") + 1]
        source = Path(argv[-1])
        calls.append({"argv": argv, "outdir": outdir, "kwargs": kwargs})
        if produce_output:
            (outdir / f"{source.stem}.{target_format}").write_bytes(output_bytes)
        import subprocess

        return subprocess.CompletedProcess(argv, returncode, stdout=b"", stderr=b"")

    monkeypatch.setattr(convert_module.shutil, "which", _which)
    monkeypatch.setattr(convert_module.subprocess, "run", _run)
    return calls


def _xlsx_bytes(sheets: dict[str, list[list[object]]]) -> bytes:
    """Build a real, tiny multi-sheet ``.xlsx`` in memory with openpyxl (MIT
    — already a transitive dependency via ``markitdown[all]``, never a new
    one). Used by the rescue-chain CSV-fallback tests, which read the
    LibreOffice-produced file with REAL openpyxl rather than a mock."""
    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(title=name)
        for row in rows:
            sheet.append(row)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("suffix, target_format", _LEGACY_OFFICE_CASES)
def test_legacy_office_suffix_is_preconverted_then_handed_to_markitdown(tmp_path, monkeypatch, suffix, target_format):
    import src.ingest.convert as convert_module

    path = tmp_path / f"legacy{suffix}"
    path.write_bytes(b"legacy office bytes")
    calls = _stub_soffice(monkeypatch, convert_module)

    markitdown_calls = []

    def _fake_markitdown(converted_path, filename, *, file_extension=None):
        markitdown_calls.append((converted_path, filename))
        return "converted text from libreoffice output"

    monkeypatch.setattr(convert_module, "_convert_markitdown", _fake_markitdown)

    result = convert_to_markdown(path, "application/octet-stream")

    assert result.engine == "libreoffice+markitdown"
    assert result.markdown == "converted text from libreoffice output"

    assert len(calls) == 1
    argv = calls[0]["argv"]
    assert argv[0] == "soffice"
    assert "--headless" in argv
    assert "--norestore" in argv
    assert argv[argv.index("--convert-to") + 1] == target_format
    assert argv[-1] == str(path)

    # markitdown ran against the LibreOffice OUTPUT, not the original file,
    # and the reported filename is still the original (source) name.
    assert len(markitdown_calls) == 1
    converted_path, filename = markitdown_calls[0]
    assert converted_path.suffix == f".{target_format}"
    assert filename == path.name

    # the temp outdir is always cleaned up
    assert not calls[0]["outdir"].exists()


def test_legacy_office_temp_dir_is_removed_even_on_failure(tmp_path, monkeypatch):
    import src.ingest.convert as convert_module

    path = tmp_path / "legacy.doc"
    path.write_bytes(b"legacy office bytes")
    calls = _stub_soffice(monkeypatch, convert_module, returncode=1)

    with pytest.raises(ConversionError):
        convert_to_markdown(path, "application/msword")

    assert len(calls) == 1
    assert not calls[0]["outdir"].exists()


def test_missing_soffice_raises_missing_conversion_dependency(tmp_path, monkeypatch):
    import src.ingest.convert as convert_module

    path = tmp_path / "legacy.doc"
    path.write_bytes(b"legacy office bytes")
    monkeypatch.setattr(convert_module.shutil, "which", lambda cmd: None)

    with pytest.raises(MissingConversionDependency) as excinfo:
        convert_to_markdown(path, "application/msword")

    assert excinfo.value.filename == "legacy.doc"
    assert excinfo.value.package == "libreoffice"
    assert excinfo.value.engine == "libreoffice+markitdown"
    assert isinstance(excinfo.value, ConversionError)


def test_libreoffice_non_zero_exit_raises_conversion_error(tmp_path, monkeypatch):
    import src.ingest.convert as convert_module

    path = tmp_path / "legacy.doc"
    path.write_bytes(b"legacy office bytes")
    _stub_soffice(monkeypatch, convert_module, returncode=1, produce_output=False)

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/msword")

    assert excinfo.value.filename == "legacy.doc"
    assert excinfo.value.engine == "libreoffice+markitdown"
    assert not isinstance(excinfo.value, MissingConversionDependency)


def test_libreoffice_timeout_raises_conversion_error(tmp_path, monkeypatch):
    import subprocess

    import src.ingest.convert as convert_module

    path = tmp_path / "legacy.ppt"
    path.write_bytes(b"legacy office bytes")
    _stub_soffice(
        monkeypatch,
        convert_module,
        side_effect=subprocess.TimeoutExpired(cmd="soffice", timeout=120),
    )

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/vnd.ms-powerpoint")

    assert excinfo.value.filename == "legacy.ppt"
    assert excinfo.value.engine == "libreoffice+markitdown"
    assert "timed out" in str(excinfo.value)


def test_libreoffice_success_with_no_output_file_raises_conversion_error(tmp_path, monkeypatch):
    import src.ingest.convert as convert_module

    path = tmp_path / "legacy.xls"
    path.write_bytes(b"legacy office bytes")
    _stub_soffice(monkeypatch, convert_module, returncode=0, produce_output=False)

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/vnd.ms-excel")

    assert excinfo.value.filename == "legacy.xls"
    assert excinfo.value.engine == "libreoffice+markitdown"


def test_legacy_office_suffixes_are_exported_and_exact(monkeypatch):
    import src.ingest.convert as convert_module

    assert convert_module.LEGACY_OFFICE_SUFFIXES == frozenset(
        {".doc", ".rtf", ".odt", ".ppt", ".odp", ".xls", ".ods", ".xlsb", ".xlsm"}
    )
    # never overlaps with the routes that already have their own engine
    assert not convert_module.LEGACY_OFFICE_SUFFIXES & convert_module.PASSTHROUGH_SUFFIXES
    assert ".pdf" not in convert_module.LEGACY_OFFICE_SUFFIXES
    # both target the same OOXML sibling a plain .xls already does
    assert convert_module.LEGACY_OFFICE_TARGETS[".xlsb"] == "xlsx"
    assert convert_module.LEGACY_OFFICE_TARGETS[".xlsm"] == "xlsx"


def test_missing_soffice_for_xlsb_raises_missing_conversion_dependency_not_a_crash(tmp_path, monkeypatch):
    """.xlsb (openpyxl cannot read the binary container at all) must fail the
    SAME attributable, retriable way the pre-2007 legacy formats already do
    — never an uncaught crash — when LibreOffice is unavailable."""
    import src.ingest.convert as convert_module

    path = tmp_path / "workbook.xlsb"
    path.write_bytes(b"xlsb binary bytes")
    monkeypatch.setattr(convert_module.shutil, "which", lambda cmd: None)

    with pytest.raises(MissingConversionDependency) as excinfo:
        convert_to_markdown(path, "application/vnd.ms-excel.sheet.binary.macroenabled.12")

    assert excinfo.value.filename == "workbook.xlsb"
    assert excinfo.value.package == "libreoffice"
    assert excinfo.value.engine == "libreoffice+markitdown"
    assert isinstance(excinfo.value, ConversionError)


def test_xlsb_libreoffice_failure_is_a_conversion_error_not_a_crash(tmp_path, monkeypatch):
    """A non-zero LibreOffice exit on `.xlsb` is an ordinary, attributable
    `ConversionError` (the crawler's `convert_failed`) — never a crash."""
    import src.ingest.convert as convert_module

    path = tmp_path / "workbook.xlsb"
    path.write_bytes(b"xlsb binary bytes")
    _stub_soffice(monkeypatch, convert_module, returncode=1, produce_output=False)

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/vnd.ms-excel.sheet.binary.macroenabled.12")

    assert excinfo.value.filename == "workbook.xlsb"
    assert excinfo.value.engine == "libreoffice+markitdown"
    assert not isinstance(excinfo.value, MissingConversionDependency)


def test_other_suffixes_are_untouched_by_the_legacy_office_route(tmp_path):
    """.docx (the modern, already-supported sibling) must keep going straight
    to markitdown — the legacy pre-conversion step is additive, never a
    detour for a format markitdown already reads natively."""
    path = _write_docx(tmp_path / "modern.docx")

    result = convert_to_markdown(
        path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert result.engine == "markitdown"


# ------------------------------------------------------------ rescue chain


def test_xlsx_rescue_chain_resaves_and_retries_once_on_markitdown_failure(tmp_path, monkeypatch):
    """Rung 1: a genuine markitdown failure on a plain ``.xlsx`` (openpyxl
    rejects it — the live finding this whole rescue chain exists for) is
    rescued by a LibreOffice resave-into-xlsx and ONE retry."""
    import src.ingest.convert as convert_module

    path = tmp_path / "report.xlsx"
    path.write_bytes(_ooxml_shaped_bytes(b"xlsx bytes markitdown rejects on the first try"))
    calls = _stub_soffice(monkeypatch, convert_module)

    attempts: list[Path] = []

    def _fake_markitdown(p, filename, *, file_extension=None):
        attempts.append(Path(p))
        if len(attempts) == 1:
            raise ConversionError(filename, "FileConversionException", engine="markitdown")
        return "recovered via libreoffice resave"

    monkeypatch.setattr(convert_module, "_convert_markitdown", _fake_markitdown)

    result = convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    assert result.engine == "libreoffice_resave+markitdown"
    assert result.rescue == "libreoffice_resave"
    assert result.markdown == "recovered via libreoffice resave"
    # markitdown ran twice: once on the original, once on the resaved copy
    assert len(attempts) == 2
    assert attempts[0] == path
    assert attempts[1] != path and attempts[1].suffix == ".xlsx"
    # exactly one soffice invocation, resaving to the SAME format (xlsx)
    assert len(calls) == 1
    argv = calls[0]["argv"]
    assert argv[argv.index("--convert-to") + 1] == "xlsx"
    assert not calls[0]["outdir"].exists()


def test_xlsx_rescue_chain_falls_back_to_csv_when_resave_retry_also_fails(tmp_path, monkeypatch):
    """Rung 2: when the resave-and-retry ALSO fails, a spreadsheet falls
    back to a LibreOffice-produced ``.xlsx`` read per-sheet as CSV, headed
    by the sheet name — never ``soffice --convert-to csv`` directly (that
    would only export the active sheet)."""
    import src.ingest.convert as convert_module

    path = tmp_path / "report.xlsx"
    path.write_bytes(_ooxml_shaped_bytes(b"xlsx bytes markitdown always rejects"))
    fallback_bytes = _xlsx_bytes({"Summary": [["Region", "Revenue"], ["EMEA", 100]], "Detail": [["Line"], ["one"]]})
    calls = _stub_soffice(monkeypatch, convert_module, output_bytes=fallback_bytes)

    def _always_fails(p, filename, *, file_extension=None):
        raise ConversionError(filename, "FileConversionException", engine="markitdown")

    monkeypatch.setattr(convert_module, "_convert_markitdown", _always_fails)

    result = convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    assert result.engine == "libreoffice_csv_fallback"
    assert result.rescue == "csv_fallback"
    assert "## Summary" in result.markdown
    assert "EMEA,100" in result.markdown
    assert "## Detail" in result.markdown
    assert "one" in result.markdown
    # two soffice invocations: rung 1's resave-and-retry, then rung 2's own
    # resave for the csv read — both temp dirs cleaned up
    assert len(calls) == 2
    for call in calls:
        assert not call["outdir"].exists()


def test_pptx_rescue_chain_falls_back_to_pdf_when_resave_retry_also_fails(tmp_path, monkeypatch):
    """Decks/documents fall back to a LibreOffice-produced PDF, run through
    this module's own PDF route (`pdf_structure`/pypdfium2) rather than a
    second text extractor."""
    import src.ingest.convert as convert_module

    path = tmp_path / "deck.pptx"
    path.write_bytes(_ooxml_shaped_bytes(b"pptx bytes markitdown always rejects"))
    pdf_bytes = _build_pdf([[("Quarterly results", 72, 700)]])
    calls = _stub_soffice(monkeypatch, convert_module, output_bytes=pdf_bytes)

    def _always_fails(p, filename, *, file_extension=None):
        raise ConversionError(filename, "FileConversionException", engine="markitdown")

    monkeypatch.setattr(convert_module, "_convert_markitdown", _always_fails)

    result = convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.presentationml.presentation")

    assert result.engine == "libreoffice_pdf_fallback"
    assert result.rescue == "pdf_fallback"
    assert "Quarterly results" in result.markdown
    assert len(calls) == 2
    assert calls[0]["argv"][calls[0]["argv"].index("--convert-to") + 1] == "pptx"
    assert calls[1]["argv"][calls[1]["argv"].index("--convert-to") + 1] == "pdf"


def test_xlsm_already_resaved_skips_rung_one_and_reuses_the_same_file_for_csv_fallback(tmp_path, monkeypatch):
    """``.xlsm`` already goes through ONE LibreOffice resave via the legacy-
    office route before markitdown ever sees it — if THAT markitdown
    attempt fails too, the rescue chain must not pay for a second, redundant
    resave-and-retry: it escalates straight to the csv fallback rung, which
    reuses the SAME already-resaved file rather than resaving again."""
    import src.ingest.convert as convert_module

    path = tmp_path / "workbook.xlsm"
    path.write_bytes(b"xlsm bytes markitdown rejects even after the legacy resave")
    resaved_bytes = _xlsx_bytes({"Data": [["A", "B"], [1, 2]]})
    calls = _stub_soffice(monkeypatch, convert_module, output_bytes=resaved_bytes)

    def _always_fails(p, filename, *, file_extension=None):
        raise ConversionError(filename, "FileConversionException", engine="markitdown")

    monkeypatch.setattr(convert_module, "_convert_markitdown", _always_fails)

    result = convert_to_markdown(path, "application/vnd.ms-excel.sheet.macroEnabled.12")

    assert result.engine == "libreoffice_csv_fallback"
    assert result.rescue == "csv_fallback"
    assert "## Data" in result.markdown
    assert "1,2" in result.markdown
    # ONE soffice call only — the legacy pre-convert to xlsx; the csv
    # fallback reads that SAME resaved file rather than a second round trip
    assert len(calls) == 1


def test_rescue_chain_reports_every_rungs_last_error_when_both_are_reached(tmp_path, monkeypatch):
    """When rung 1's OWN re-save succeeds but markitdown still rejects the
    resaved copy, rung 2 is reached and the raised `ConversionError` names
    EACH rung's own last error — not just markitdown's — so the next
    reconciliation pass can tell WHICH step to fix."""
    import src.ingest.convert as convert_module

    path = tmp_path / "report.xlsx"
    path.write_bytes(_ooxml_shaped_bytes(b"xlsx bytes"))
    # Every soffice invocation SUCCEEDS (exit 0, produces a file) — a
    # resave-itself failure is a DIFFERENT scenario (see the "unopenable"
    # test below) that stops the chain before rung 2 is even reached. Here
    # both soffice calls succeed structurally; it is the SUBSEQUENT
    # markitdown/openpyxl parse of what they produced that fails each time,
    # which is what lets BOTH rungs actually run.
    _stub_soffice(monkeypatch, convert_module, output_bytes=b"not a real xlsx")

    def _always_fails(p, filename, *, file_extension=None):
        raise ConversionError(filename, "FileConversionException: bad zip", engine="markitdown")

    monkeypatch.setattr(convert_module, "_convert_markitdown", _always_fails)

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    message = str(excinfo.value)
    assert "markitdown: " in message and "bad zip" in message
    assert "libreoffice resave+retry: " in message
    assert "csv fallback: " in message
    assert excinfo.value.filename == "report.xlsx"


def test_rescue_chain_stops_after_rung_one_when_the_resave_itself_is_unopenable(tmp_path, monkeypatch):
    """2026-09-04 finding #66 item 3 (live finding: hundreds of doomed
    spreadsheets held 16-22 conversion children for minutes each). When
    rung 1's OWN LibreOffice re-save fails because the SOURCE file itself is
    corrupt/unopenable (a non-zero, non-signal exit), rung 2's CSV/PDF
    fallback — which would reach the IDENTICAL LibreOffice mechanism on the
    IDENTICAL bytes — is skipped rather than retried."""
    import src.ingest.convert as convert_module

    path = tmp_path / "report.xlsx"
    path.write_bytes(_ooxml_shaped_bytes(b"xlsx bytes"))
    calls = _stub_soffice(monkeypatch, convert_module, returncode=1, produce_output=False)

    def _always_fails(p, filename, *, file_extension=None):
        raise ConversionError(filename, "FileConversionException: bad zip", engine="markitdown")

    monkeypatch.setattr(convert_module, "_convert_markitdown", _always_fails)

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    message = str(excinfo.value)
    assert "markitdown: " in message and "bad zip" in message
    assert "libreoffice resave: " in message
    assert "skipped" in message
    assert message.count("exited with status 1") == 1  # rung 1's resave only — rung 2 never ran
    assert excinfo.value.filename == "report.xlsx"
    assert excinfo.value.error_class == convert_module.ERROR_CLASS_LIBREOFFICE_NO_OUTPUT
    # exactly ONE soffice call — the failed resave; rung 2 never shells out again
    assert len(calls) == 1


def test_rescue_chain_stops_after_rung_one_on_a_libreoffice_timeout(tmp_path, monkeypatch):
    """A timeout on rung 1's own re-save also stops the chain — retrying a
    slower rung after LibreOffice already timed out on this file would only
    time out again (2026-09-04 finding #66 item 3)."""
    import subprocess

    import src.ingest.convert as convert_module

    path = tmp_path / "report.xlsx"
    path.write_bytes(_ooxml_shaped_bytes(b"xlsx bytes"))
    _stub_soffice(monkeypatch, convert_module, side_effect=subprocess.TimeoutExpired(cmd="soffice", timeout=120))
    run_count = 0
    real_run = convert_module.subprocess.run

    def _counting_run(*args, **kwargs):
        nonlocal run_count
        run_count += 1
        return real_run(*args, **kwargs)

    monkeypatch.setattr(convert_module.subprocess, "run", _counting_run)

    def _always_fails(p, filename, *, file_extension=None):
        raise ConversionError(filename, "FileConversionException", engine="markitdown")

    monkeypatch.setattr(convert_module, "_convert_markitdown", _always_fails)

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    assert excinfo.value.error_class == convert_module.ERROR_CLASS_TIMEOUT
    assert run_count == 1


@pytest.mark.parametrize(
    "returncode, expected_class",
    [(-9, "memory_kill"), (-11, "worker_crash")],
)
def test_rescue_chain_stops_after_rung_one_on_a_signal_killed_resave(tmp_path, monkeypatch, returncode, expected_class):
    """A resave killed by a SIGNAL (the LibreOffice subprocess itself was
    killed — by a memory guard/OOM killer for SIGKILL, or crashed for any
    other signal) is environmental, not a property of the file — but it
    STILL stops the chain: retrying the identical subprocess mechanism on
    the identical bytes right away is not a rescue."""
    import src.ingest.convert as convert_module

    path = tmp_path / "report.xlsx"
    path.write_bytes(_ooxml_shaped_bytes(b"xlsx bytes"))
    calls = _stub_soffice(monkeypatch, convert_module, returncode=returncode, produce_output=False)

    def _always_fails(p, filename, *, file_extension=None):
        raise ConversionError(filename, "FileConversionException", engine="markitdown")

    monkeypatch.setattr(convert_module, "_convert_markitdown", _always_fails)

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    assert excinfo.value.error_class == expected_class
    assert len(calls) == 1


def test_rescue_chain_never_fires_for_a_non_rescuable_suffix(tmp_path, monkeypatch):
    """A format outside the rescue chain's suffix maps (``.html``, here)
    must never even PROBE for LibreOffice on a markitdown failure — the
    ordinary, immediate `ConversionError` is unchanged."""
    import src.ingest.convert as convert_module

    path = tmp_path / "page.html"
    path.write_text("<html></html>", encoding="utf-8")

    def _boom(p, filename, *, file_extension=None):
        raise ConversionError(filename, "markitdown could not convert this file", engine="markitdown")

    monkeypatch.setattr(convert_module, "_convert_markitdown", _boom)
    which_calls: list[str] = []
    monkeypatch.setattr(convert_module.shutil, "which", lambda cmd: which_calls.append(cmd) or None)

    with pytest.raises(ConversionError):
        convert_to_markdown(path, "text/html")

    assert which_calls == [], "a non-rescuable suffix must never even check for libreoffice"


# ---------------------------------------------------- size-scaled conversion budget


def test_conversion_budget_seconds_scales_with_input_size():
    from src.ingest.convert import CONVERSION_BUDGET_BASE_SECONDS, conversion_budget_seconds

    assert conversion_budget_seconds(0) == CONVERSION_BUDGET_BASE_SECONDS
    fifteen_mb = 15 * 1024 * 1024
    # the live finding this constant is sized against: 221 large xlsx/xlsm
    # files averaging 15 MB hit the flat 300s budget
    assert conversion_budget_seconds(fifteen_mb) == pytest.approx(300.0 + 15 * 20.0)


def test_conversion_budget_seconds_caps_at_the_ceiling():
    from src.ingest.convert import CONVERSION_BUDGET_MAX_SECONDS, conversion_budget_seconds

    huge = 500 * 1024 * 1024
    assert conversion_budget_seconds(huge) == CONVERSION_BUDGET_MAX_SECONDS


def test_conversion_budget_seconds_zero_base_disables_it():
    from src.ingest.convert import conversion_budget_seconds

    assert conversion_budget_seconds(15 * 1024 * 1024, base_seconds=0) == 0.0


def test_conversion_budget_seconds_respects_a_custom_base():
    from src.ingest.convert import conversion_budget_seconds

    assert conversion_budget_seconds(1024 * 1024, base_seconds=120) == pytest.approx(140.0)


# ------------------------------------------------------------- xlsx streaming


def test_large_xlsx_routes_to_openpyxl_streaming_not_markitdown(tmp_path, monkeypatch):
    import src.ingest.convert as convert_module

    path = tmp_path / "huge.xlsx"
    path.write_bytes(_xlsx_bytes({"Sheet1": [["a", "b"], [1, 2]]}))
    # force the size-threshold branch regardless of this tiny fixture's real size
    monkeypatch.setattr(convert_module, "LARGE_XLSX_STREAMING_THRESHOLD_BYTES", 0)
    markitdown_called: list[int] = []
    monkeypatch.setattr(
        convert_module, "_convert_markitdown", lambda p, f, **_kw: markitdown_called.append(1) or "nope"
    )

    result = convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    assert result.engine == "openpyxl_streaming"
    assert result.rescue == ""
    assert "## Sheet1" in result.markdown
    assert "1,2" in result.markdown
    assert markitdown_called == [], "the whole point of the streaming route is to skip markitdown"


def test_large_xlsx_streaming_stops_early_on_a_50k_row_workbook(tmp_path, monkeypatch):
    """The literal live finding this route exists for: a huge workbook must
    stop reading once the document-wide char cap is reached, not read every
    row and truncate afterward — proven here by timing (streaming 50k rows
    down to a couple thousand characters must stay fast) and by the output
    staying far short of what reading all 50k rows would produce."""
    import time

    import openpyxl

    import src.ingest.convert as convert_module

    path = tmp_path / "wide.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Rows"
    row = ["x" * 50] * 10
    for _ in range(50_000):
        sheet.append(row)
    workbook.save(str(path))

    monkeypatch.setattr(convert_module, "LARGE_XLSX_STREAMING_THRESHOLD_BYTES", 0)

    started = time.monotonic()
    result = convert_to_markdown(
        path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", max_chars=2_000
    )
    elapsed = time.monotonic() - started

    assert result.engine == "openpyxl_streaming"
    assert elapsed < 5.0, f"streaming a 50k-row workbook under a small cap took {elapsed:.2f}s"
    # 50k rows of this shape would produce well over 2.5M raw characters if
    # read in full — this stays close to the cap, proof the READ stopped
    # early rather than merely the OUTPUT being cut late
    assert len(result.markdown) < 50_000


def test_large_xlsx_streaming_falls_back_to_libreoffice_resave_on_openpyxl_failure(tmp_path, monkeypatch):
    """If openpyxl itself cannot open the large file directly, a fresh
    LibreOffice resave (a clean, re-encoded copy) is tried before giving
    up — the same CSV-fallback mechanism the rescue chain's rung 2 uses,
    reached directly here since markitdown is never attempted on this
    route."""
    import src.ingest.convert as convert_module

    path = tmp_path / "huge.xlsx"
    path.write_bytes(_ooxml_shaped_bytes(b"not a real xlsx, openpyxl will refuse it"))
    monkeypatch.setattr(convert_module, "LARGE_XLSX_STREAMING_THRESHOLD_BYTES", 0)
    resaved_bytes = _xlsx_bytes({"Recovered": [["ok"], [1]]})
    calls = _stub_soffice(monkeypatch, convert_module, output_bytes=resaved_bytes)

    result = convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    assert result.engine == "libreoffice_csv_fallback"
    assert result.rescue == "csv_fallback"
    assert "## Recovered" in result.markdown
    assert len(calls) == 1
    assert calls[0]["argv"][calls[0]["argv"].index("--convert-to") + 1] == "xlsx"


# ------------------------------------------------------------------ failures


def test_corrupt_pdf_raises_conversion_error_naming_the_file(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4\nthis is not a pdf body at all\n")

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/pdf")

    assert excinfo.value.filename == "broken.pdf"
    assert "broken.pdf" in str(excinfo.value)
    assert excinfo.value.engine == "pypdfium2"


def test_corrupt_archive_format_raises_conversion_error_naming_the_file(tmp_path):
    # A truncated EPUB package: the zip header is there, the archive is not.
    # markitdown surfaces this as its own FileConversionException wrapping a
    # zipfile.BadZipFile — one of many backend exception types this module
    # deliberately funnels into ConversionError. EPUB rather than an office
    # suffix on purpose: the OOXML archive guard would refuse a .docx/.xlsx
    # before markitdown ever saw it, and this test is about the funnel.
    path = tmp_path / "broken.epub"
    path.write_bytes(b"PK\x03\x04\x00\x00truncated archive")

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/epub+zip")

    assert excinfo.value.filename == "broken.epub"
    assert excinfo.value.engine == "markitdown"


def test_office_file_that_is_not_an_archive_is_a_conversion_error(tmp_path):
    """A ``.docx``/``.pptx``/``.xlsx`` is an OOXML package — a zip — or it is
    not that file type at all. markitdown sniffs content and would read ASCII
    bytes behind an office suffix as prose (in practice as UTF-16 mojibake),
    silently indexing garbage; the converter refuses before it gets there."""
    path = tmp_path / "notes.docx"
    path.write_bytes(b"PK\x03\x04 this is not a zip archive at all")

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/octet-stream")

    assert not isinstance(excinfo.value, MissingConversionDependency)
    assert excinfo.value.filename == "notes.docx"
    assert "archive" in str(excinfo.value)


def test_missing_file_raises_conversion_error(tmp_path):
    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(tmp_path / "gone.docx", "application/octet-stream")

    assert excinfo.value.filename == "gone.docx"


@pytest.mark.parametrize("suffix", [".pbix", ".one"])
def test_a_format_with_no_registered_backend_raises_unsupported_not_a_plain_conversion_error(tmp_path, suffix):
    # Power BI (.pbix) and OneNote (.one): no markitdown converter's
    # `accepts()` matches these at all (verified empirically — see
    # `UnsupportedConversionFormat`'s docstring), so markitdown raises its
    # own `UnsupportedFormatException` before any backend is even tried.
    # Random bytes, not a crafted file: the point is that NOTHING attempts
    # this format, regardless of content.
    path = tmp_path / f"deck{suffix}"
    path.write_bytes(bytes((i * 37) % 256 for i in range(2048)))

    with pytest.raises(UnsupportedConversionFormat) as excinfo:
        convert_to_markdown(path, "application/octet-stream")

    assert isinstance(excinfo.value, ConversionError)
    assert excinfo.value.filename == f"deck{suffix}"
    assert excinfo.value.engine == "markitdown"


def test_a_directory_is_not_convertible(tmp_path):
    directory = tmp_path / "folder.pdf"
    directory.mkdir()

    with pytest.raises(ConversionError):
        convert_to_markdown(directory, "application/pdf")


def test_conversion_error_never_kills_a_crawl_loop(tmp_path):
    """The contract the crawler relies on: one bad file, one caught error."""
    good = tmp_path / "good.txt"
    good.write_text("fine", encoding="utf-8")
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"not a pdf")

    statuses = []
    for path in (good, bad, good):
        try:
            statuses.append(convert_to_markdown(path, "application/octet-stream").engine)
        except ConversionError as exc:
            statuses.append(f"error:{exc.filename}")

    assert statuses == ["passthrough", "error:bad.pdf", "passthrough"]


# ------------------------------------------------------- missing dependencies


@pytest.mark.parametrize(
    "module, name, mime",
    [
        ("markitdown", "handbook.docx", "application/octet-stream"),
        ("pypdfium2", "report.pdf", "application/pdf"),
    ],
)
def test_missing_backend_raises_a_typed_error_naming_the_extra(tmp_path, monkeypatch, module, name, mime):
    # ``None`` in sys.modules makes ``from <module> import X`` raise ImportError,
    # which is exactly what an uninstalled extra looks like.
    monkeypatch.setitem(sys.modules, module, None)
    path = tmp_path / name
    if name.endswith(".docx"):
        # The OOXML archive guard runs before any engine is imported, so the
        # office fixture has to be a zip for the dependency probe to be the
        # thing that fails.
        with zipfile.ZipFile(path, "w") as package:
            package.writestr("[Content_Types].xml", "<Types/>")
    else:
        path.write_bytes(b"irrelevant, the import fails first")

    with pytest.raises(MissingConversionDependency) as excinfo:
        convert_to_markdown(path, mime)

    assert excinfo.value.package == module
    assert "extraction" in str(excinfo.value)
    assert isinstance(excinfo.value, ConversionError)


class _RaisingMetaPathFinder:
    """A ``sys.meta_path`` entry that raises ``exc`` instead of resolving
    ``module_name`` — simulates a lazy import failing with an arbitrary
    exception type (not just ``ImportError``), for both an ``import x`` and
    a ``from x import y`` statement, since it intercepts before ``x`` is
    even bound into ``sys.modules``."""

    def __init__(self, module_name: str, exc: BaseException) -> None:
        self._module_name = module_name
        self._exc = exc

    def find_spec(self, fullname, path, target=None):
        if fullname == self._module_name:
            raise self._exc


@pytest.mark.parametrize(
    "module, name, mime",
    [
        ("markitdown", "handbook.docx", "application/octet-stream"),
        ("pypdfium2", "report.pdf", "application/pdf"),
    ],
)
def test_missing_backend_import_failure_carries_the_causes_type_and_text(tmp_path, monkeypatch, module, name, mime):
    """A live finding: a MemoryError surfacing from deep inside a lazy
    ``import markitdown``/``import pypdfium2`` (a conversion child's own
    numpy/OpenBLAS init failing under its RLIMIT_AS) must not read as a
    plain "not installed" — that cost 30 minutes to diagnose. The message
    must name the real exception type and carry its text."""
    monkeypatch.delitem(sys.modules, module, raising=False)
    finder = _RaisingMetaPathFinder(
        module, MemoryError("OpenBLAS error: Memory allocation still failed after 10 retries.")
    )
    monkeypatch.setattr(sys, "meta_path", [finder, *sys.meta_path])
    path = tmp_path / name
    # A zip for the .docx case, for the same reason as the parametrized test
    # above: the archive guard runs before any engine is imported.
    path.write_bytes(_ooxml_shaped_bytes() if name.endswith(".docx") else b"irrelevant, the import fails first")

    with pytest.raises(MissingConversionDependency) as excinfo:
        convert_to_markdown(path, mime)

    assert excinfo.value.package == module
    message = str(excinfo.value)
    assert "MemoryError" in message
    assert "OpenBLAS error: Memory allocation still failed after 10 retries." in message
    assert "conversion child's memory limit" in message
    assert isinstance(excinfo.value, ConversionError)


def test_missing_backend_import_failure_message_truncates_a_long_cause(tmp_path, monkeypatch):
    """The cause's text is capped at 200 chars so one pathological exception
    message can't blow up the crawl's error log/report."""
    monkeypatch.delitem(sys.modules, "markitdown", raising=False)
    long_text = "x" * 500
    finder = _RaisingMetaPathFinder("markitdown", RuntimeError(long_text))
    monkeypatch.setattr(sys, "meta_path", [finder, *sys.meta_path])
    path = tmp_path / "handbook.docx"
    path.write_bytes(_ooxml_shaped_bytes())

    with pytest.raises(MissingConversionDependency) as excinfo:
        convert_to_markdown(path, "application/octet-stream")

    message = str(excinfo.value)
    assert ("x" * 200) in message
    assert ("x" * 201) not in message


def test_module_imports_without_the_extraction_extra(monkeypatch):
    """Import-time cost is zero: the backends are only imported on demand."""
    import importlib

    import src.ingest as ingest_package

    monkeypatch.setitem(sys.modules, "markitdown", None)
    monkeypatch.setitem(sys.modules, "pypdfium2", None)
    # The re-import below rebinds the PACKAGE attribute to a second module
    # object; ``sys.modules`` is restored by ``delitem`` but the attribute is
    # not, and a later ``from src.ingest import convert`` would then patch a
    # module nothing else in the process calls. Pin it so teardown restores it.
    monkeypatch.setattr(ingest_package, "convert", ingest_package.convert, raising=False)
    monkeypatch.delitem(sys.modules, "src.ingest.convert", raising=False)

    module = importlib.import_module("src.ingest.convert")

    assert module.convert_to_markdown is not None


# ------------------------------------------------------------ docling engine


@pytest.fixture
def fake_docling(monkeypatch):
    """A stand-in for the ``[docling]`` extra.

    The real one pulls torch and cannot be installed in the default test
    environment; the ``rich-extras`` CI job covers it for real. What is under
    test here is the ROUTING — which engine answers, and what happens when
    Docling is present but fails — so the stand-in only has to look like
    ``docling.document_converter.DocumentConverter`` from the call site's
    side: ``.convert(path).document.export_to_markdown()``.

    ``docling_capability`` is patched alongside: it is an import-spec probe,
    and a synthetic module in ``sys.modules`` has no spec to find.
    """
    from src.ingest import convert

    calls: list[Path] = []
    state = {"raise": False, "markdown": "# From docling\n\nconverted by the stand-in"}

    class _Document:
        def export_to_markdown(self) -> str:
            return state["markdown"]

    class _Result:
        document = _Document()

    class DocumentConverter:
        def convert(self, path):
            calls.append(Path(path))
            if state["raise"]:
                raise RuntimeError("docling choked on purpose")
            return _Result()

    package = types.ModuleType("docling")
    module = types.ModuleType("docling.document_converter")
    module.DocumentConverter = DocumentConverter  # type: ignore[attr-defined]
    package.document_converter = module  # type: ignore[attr-defined]
    monkeypatch.setitem(sys.modules, "docling", package)
    monkeypatch.setitem(sys.modules, "docling.document_converter", module)
    monkeypatch.setattr(convert, "docling_capability", lambda: True)
    return types.SimpleNamespace(calls=calls, state=state)


def test_office_document_prefers_docling_when_installed(tmp_path, fake_docling):
    """The rich image's whole point: layout-aware parsing takes precedence
    over markitdown for office documents when the extra is there."""
    path = _write_docx(tmp_path / "handbook.docx")

    result = convert_to_markdown(path, "application/octet-stream")

    assert result.engine == "docling"
    assert "converted by the stand-in" in result.markdown
    assert fake_docling.calls == [path]


def test_docling_failure_falls_through_to_markitdown(tmp_path, fake_docling):
    """Docling present but choking on THIS document is not a rejection:
    markitdown gets the file next, exactly as the Collections extractor
    always did when Docling failed."""
    fake_docling.state["raise"] = True
    path = _write_docx(tmp_path / "handbook.docx")

    result = convert_to_markdown(path, "application/octet-stream")

    assert result.engine == "markitdown"
    assert "Agnes Handbook" in result.markdown


def test_docling_failure_without_markitdown_blames_the_file_not_a_missing_extra(tmp_path, fake_docling, monkeypatch):
    """On a rich image built without ``[extraction]`` a document Docling
    cannot read has NO second reader — but that is the file's problem, and
    ``MissingConversionDependency`` would send an operator to install an extra
    for a document that would still fail."""
    fake_docling.state["raise"] = True
    monkeypatch.setitem(sys.modules, "markitdown", None)
    path = _write_docx(tmp_path / "handbook.docx")

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/octet-stream")

    assert not isinstance(excinfo.value, MissingConversionDependency)
    assert excinfo.value.engine == "docling"
    assert excinfo.value.filename == "handbook.docx"


def test_pdf_never_routes_to_docling(tmp_path, fake_docling):
    """Owner decision 2026-08-31 — one PDF pipeline, no dual modes: the
    structure pass is the PDF route on every image, Docling or not."""
    path = tmp_path / "hello.pdf"
    path.write_bytes(_build_pdf([[("Hello Agnes", 72, 700), ("Second line", 72, 660)]]))

    result = convert_to_markdown(path, "application/pdf")

    assert result.engine == "pypdfium2"
    assert fake_docling.calls == []


def test_passthrough_never_routes_to_docling(tmp_path, fake_docling):
    path = tmp_path / "notes.md"
    path.write_text("# already markdown\n", encoding="utf-8")

    result = convert_to_markdown(path, "text/markdown")

    assert result.engine == "passthrough"
    assert fake_docling.calls == []


# ------------------------------------------------------------- suffix hint


def test_suffix_hint_routes_a_file_stored_without_its_extension(tmp_path):
    """Collections store an upload as ``<sha256><ext>`` and keep the declared
    type as a column; a caller that knows the type better than the storage
    name says so, and routing follows the hint rather than the path."""
    blob = tmp_path / "3f2a9c0e"
    blob.write_text("# stored under a hash\n\nbody", encoding="utf-8")

    assert convert_to_markdown(blob, "application/octet-stream", suffix=".md").engine == "passthrough"

    office = _write_docx(tmp_path / "7b1d")
    result = convert_to_markdown(office, "application/octet-stream", suffix=".docx")
    assert result.engine == "markitdown"
    assert "Agnes Handbook" in result.markdown


# --------------------------------------------------------- licence invariant


#: Modules this converter may never import, whatever the quality argument for
#: them: Agnes ships under PolyForm Small Business 1.0.0 and cannot vendor
#: AGPL-3.0 code (spec §9.1). Matched on the imported module, not on the file
#: text, so the docstring can keep explaining *why* they are banned.
_AGPL_MODULES = {"fitz", "pymupdf", "pymupdf4llm", "frontend"}


def test_no_agpl_dependency_is_imported_by_the_converter():
    import ast

    source = (Path(__file__).resolve().parents[1] / "src" / "ingest" / "convert.py").read_text(encoding="utf-8")
    imported: set[str] = set()
    for node in ast.walk(ast.parse(source)):
        if isinstance(node, ast.Import):
            imported.update(alias.name.split(".")[0] for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            imported.add(node.module.split(".")[0])

    assert not imported & _AGPL_MODULES, f"AGPL dependency imported: {sorted(imported & _AGPL_MODULES)}"
    assert {"markitdown", "pypdfium2"} <= imported


def test_structure_pass_failure_raises_conversion_error(tmp_path, monkeypatch):
    """``pdf_structure.reconstruct_pdf`` is the ONLY PDF route (owner
    decision 2026-08-31 — one pipeline, no dual modes). It already degrades
    per page inside itself, so an exception escaping it means the FILE is not
    convertible: that surfaces as a ``ConversionError`` naming the file, which
    the crawler counts in ``convert_failed`` and walks past — never a silent
    second attempt through a parallel plain-text implementation."""
    from src.ingest import pdf_structure

    def _boom(path, max_pages=None):
        raise RuntimeError("structure pass broken on purpose")

    monkeypatch.setattr(pdf_structure, "reconstruct_pdf", _boom)

    path = tmp_path / "fallback.pdf"
    path.write_bytes(_build_pdf([[("Hello Agnes", 72, 700), ("Second line", 72, 660)]]))

    with pytest.raises(ConversionError) as excinfo:
        convert_to_markdown(path, "application/pdf")

    assert excinfo.value.filename == "fallback.pdf"
    assert excinfo.value.engine == "pypdfium2"


def test_legacy_office_runs_on_a_per_process_profile_and_serializes(tmp_path, monkeypatch):
    """Two headless LibreOffice instances on the SAME user profile do not
    coexist — the second exits 1 (measured live: 3 of 6 concurrent
    conversions failed). Every soffice call must therefore carry a
    ``-env:UserInstallation=`` pointing at a directory private to this
    process, reused across calls (a fresh profile costs seconds on first
    start), and calls within one process are serialized by a lock."""
    import os

    import src.ingest.convert as convert_module

    monkeypatch.setattr(convert_module, "_LIBREOFFICE_PROFILES", {})
    monkeypatch.setattr(convert_module, "_convert_markitdown", lambda p, f, **_kw: "text")
    calls = _stub_soffice(monkeypatch, convert_module)

    for name in ("a.doc", "b.xls"):
        path = tmp_path / name
        path.write_bytes(b"legacy")
        convert_to_markdown(path, "application/octet-stream")

    profiles = []
    for call in calls:
        env = [a for a in call["argv"] if a.startswith("-env:UserInstallation=file://")]
        assert len(env) == 1, call["argv"]
        profiles.append(env[0].split("file://", 1)[1])
    # same private profile for both calls of this process, and it exists
    assert profiles[0] == profiles[1]
    assert os.path.isdir(profiles[0])
    assert str(os.getpid()) in profiles[0]
    # the profile the calls used is exactly the one the helper hands out for
    # this pid (module identity resolved through the function that ran, so a
    # second import path of the same file cannot fool the check)
    import sys

    live = sys.modules[convert_to_markdown.__module__]
    assert profiles[0] == live._libreoffice_profile_dir()
    assert os.getpid() in live._LIBREOFFICE_PROFILES
    assert live._LIBREOFFICE_LOCK is not None


# ------------------------------------------------------ embedded-picture disclosure
#
# The live finding these tests pin (2026-09-09): a reader asked about content
# that lives in a picture inside a Word/PowerPoint file, and the answer was
# assembled from prose instead — without being told the picture existed. Two
# separate defects, both in markitdown's OWN image handling that this module
# had never touched before: a docx picture came back as a literal
# ``![](data:image/png;base64...)`` (markitdown's own truncation of a payload
# this module never keeps either), and a pptx picture came back as
# ``![](Picture3.jpg)`` — a name that repeats across slides and, worse,
# across two entirely unrelated documents, reading like a stable, fetchable
# filename when it is neither. These tests assert the READER-VISIBLE
# behaviour (what ends up in the indexed markdown, and the count a caller
# gets back) — never `markitdown`'s or `mammoth`'s internal call shape.


def test_docx_embedded_picture_placeholder_is_disclosed_not_base64(tmp_path):
    path = _write_docx_with_images(tmp_path / "sow.docx", [("Phase 1 Roadmap", "rId2")])

    result = convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    assert result.engine == "markitdown"
    assert result.image_count == 1
    assert "base64" not in result.markdown
    assert "data:image" not in result.markdown
    assert "[image 1 of 1 in this document — not indexed" in result.markdown
    assert 'section "Phase 1 Roadmap"' in result.markdown


def test_docx_without_images_reports_zero_image_count_and_is_unaffected(tmp_path):
    path = _write_docx(tmp_path / "handbook.docx")

    result = convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    assert result.image_count == 0
    assert "Agnes Handbook" in result.markdown
    assert "Revenue is recognised on delivery." in result.markdown
    assert "not indexed" not in result.markdown


def test_pptx_embedded_picture_placeholder_is_disclosed_not_a_collidable_filename(tmp_path):
    path = _write_pptx_with_images(tmp_path / "deck.pptx", slide_count=1)

    result = convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.presentationml.presentation")

    assert result.engine == "markitdown"
    assert result.image_count == 1
    # no PictureN.jpg-shaped name survives — the whole point is that it is no
    # longer readable as a filename that could (wrongly) be fetched or
    # cross-referenced against another document's own "PictureN".
    assert ".jpg" not in result.markdown
    assert "[image 1 of 1 in this document — not indexed, slide 1]" in result.markdown


def test_pptx_multiple_slides_number_images_and_locate_each_one(tmp_path):
    path = _write_pptx_with_images(tmp_path / "deck.pptx", slide_count=3)

    result = convert_to_markdown(path, "application/vnd.openxmlformats-officedocument.presentationml.presentation")

    assert result.image_count == 3
    assert "[image 1 of 3 in this document — not indexed, slide 1]" in result.markdown
    assert "[image 2 of 3 in this document — not indexed, slide 2]" in result.markdown
    assert "[image 3 of 3 in this document — not indexed, slide 3]" in result.markdown


def test_legacy_office_markitdown_output_also_gets_image_disclosure(tmp_path, monkeypatch):
    """The rewrite is keyed off the resulting ENGINE, not the caller's route:
    a legacy ``.doc`` resaved to docx and handed to the SAME markitdown call
    (``libreoffice+markitdown``) must be disclosed identically to a direct
    ``.docx`` upload — a reader should never learn less from a document that
    happened to need a LibreOffice resave first."""
    import src.ingest.convert as convert_module

    path = tmp_path / "legacy.doc"
    path.write_bytes(b"legacy office bytes")
    _stub_soffice(monkeypatch, convert_module)

    def _fake_markitdown(converted_path, filename, *, file_extension=None):
        return "# Q3 Notes\n\n![](data:image/png;base64...)"

    monkeypatch.setattr(convert_module, "_convert_markitdown", _fake_markitdown)

    result = convert_to_markdown(path, "application/octet-stream")

    assert result.engine == "libreoffice+markitdown"
    assert result.image_count == 1
    assert "base64" not in result.markdown
    assert 'section "Q3 Notes"' in result.markdown


def test_disclose_image_placeholders_is_a_noop_on_plain_text():
    from src.ingest.convert import _disclose_image_placeholders

    text = "just some prose with no images at all"

    rewritten, count = _disclose_image_placeholders(text)

    assert rewritten == text
    assert count == 0


def test_disclose_image_placeholders_mixed_docx_and_pptx_shapes_in_one_pass():
    """A synthetic mix of both marker shapes in one pass — exercises both
    regex branches, not a realistic single document (a real one is either
    pptx-with-jpg-names or docx-with-base64, never both). Once a slide
    marker has been seen, it wins over ANY later heading for every image
    after it, including a base64 (docx-shaped) one — see
    ``test_disclose_image_placeholders_slide_number_wins_over_a_later_heading``
    for the realistic single-slide case this generalizes."""
    from src.ingest.convert import _disclose_image_placeholders

    text = "<!-- Slide number: 1 -->\n![](Picture3.jpg)\n# A Heading\n![alt text](data:image/jpeg;base64...)\n"

    rewritten, count = _disclose_image_placeholders(text)

    assert count == 2
    assert "Picture3.jpg" not in rewritten
    assert "base64" not in rewritten
    assert "[image 1 of 2 in this document — not indexed, slide 1]" in rewritten
    assert "[image 2 of 2 in this document — not indexed, slide 1]" in rewritten
    assert "in section" not in rewritten
    # the slide/heading markers themselves are preserved verbatim — only the
    # image markdown is rewritten
    assert "<!-- Slide number: 1 -->" in rewritten
    assert "# A Heading" in rewritten


def test_disclose_image_placeholders_slide_number_wins_over_a_later_heading():
    """A pptx slide's own title becomes a ``#`` heading right after that
    slide's number comment (see the module docstring) — a titled slide must
    keep its own slide number, never fall back to "in section <title>" just
    because the heading was the more RECENT marker (live finding
    2026-09-09: this exact ordering lost the slide number)."""
    from src.ingest.convert import _disclose_image_placeholders

    text = "<!-- Slide number: 4 -->\n# Architecture\n![](Picture1.jpg)\n"

    rewritten, count = _disclose_image_placeholders(text)

    assert count == 1
    assert "[image 1 of 1 in this document — not indexed, slide 4]" in rewritten
    assert "in section" not in rewritten


def test_disclose_image_placeholders_new_slide_resets_a_stale_heading():
    """A heading tracked on one slide must not leak onto the NEXT slide's
    own images once a new slide marker has been seen — the slide number
    still wins regardless, but the reset keeps the two pieces of state
    honest independently of that priority."""
    from src.ingest.convert import _disclose_image_placeholders

    text = "<!-- Slide number: 1 -->\n# First Slide\n<!-- Slide number: 2 -->\n![](Picture1.jpg)\n"

    rewritten, count = _disclose_image_placeholders(text)

    assert count == 1
    assert "[image 1 of 1 in this document — not indexed, slide 2]" in rewritten
    assert "First Slide" not in rewritten.split("[image")[-1]


def test_disclose_image_placeholders_word_heading_untouched_by_the_slide_gate():
    """A Word document never emits a slide-number comment at all — the
    heading-tracking path (unchanged by the slide/heading split) is what a
    docx picture's location still comes from."""
    from src.ingest.convert import _disclose_image_placeholders

    text = "# Q3 Notes\n\n![](data:image/png;base64...)\n"

    rewritten, count = _disclose_image_placeholders(text)

    assert count == 1
    assert '[image 1 of 1 in this document — not indexed, in section "Q3 Notes"]' in rewritten


def test_disclose_image_placeholders_bare_jpg_name_without_a_slide_marker_is_left_alone():
    """The bare ``name.jpg`` shape is markitdown's OWN pptx lost-picture
    convention, but the identical markdown is also what an ordinary,
    un-lost HTML relative image reference converts to (see
    ``test_html_relative_image_is_not_falsely_disclosed_as_lost``). Without
    a PowerPoint slide-number marker anywhere in the document, this is not a
    lost picture — left completely unchanged, not counted (live finding
    2026-09-09)."""
    from src.ingest.convert import _disclose_image_placeholders

    text = "# Quarterly\n\n![Logo](logo.jpg)\n\nEMEA grew.\n"

    rewritten, count = _disclose_image_placeholders(text)

    assert count == 0
    assert rewritten == text


def test_disclose_image_placeholders_truncates_a_very_long_heading():
    from src.ingest.convert import _disclose_image_placeholders

    heading_text = "x" * 200
    text = f"# {heading_text}\n![](data:image/png;base64...)\n"

    rewritten, count = _disclose_image_placeholders(text)

    assert count == 1
    disclosure_line = next(line for line in rewritten.splitlines() if line.startswith("[image"))
    assert len(disclosure_line) < 150
