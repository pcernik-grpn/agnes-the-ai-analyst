"""What the optional extras change, and what the absence of them must say.

Two capabilities ship as opt-in extras because both pull torch:
``[docling]`` (layout-aware office/PDF parsing) and ``[embeddings]`` (hybrid
retrieval). Office documents do NOT depend on the first: the shared converter
(``src/ingest/convert.py``) reads ``.docx``/``.pptx`` through markitdown from
the ``[extraction]`` extra the default image ships, and prefers Docling only
when it is installed. A deployment with NEITHER parser still ACCEPTS those
uploads — they are on the upload allowlist — and then cannot index them; one
without ``[embeddings]`` ranks search results lexically only. Neither
degradation is wrong; both are invisible, which is the problem this suite pins:

* the rejection names the missing extra, so "why was my file rejected"
  has an answer that is not "your file is broken"
* the capability probes agree with what is actually installed

The extras-present direction runs only where the extras ARE installed
(``-rich`` image, and the ``rich-extras`` CI job) — skipped elsewhere rather
than mocked, because a mocked docling proves nothing about whether a real
docx parses.
"""

from __future__ import annotations

import importlib.util
import zipfile

import pytest

from src.ingest.text_extract import UnsupportedDocument, docling_capability, extract_text

_HAS_DOCLING = docling_capability()
_HAS_MARKITDOWN = importlib.util.find_spec("markitdown") is not None


def _minimal_docx(tmp_path, name="doc.docx") -> str:
    """A real (if minimal) OOXML docx — a zip with the parts a parser needs."""
    path = tmp_path / name
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.'
            'wordprocessingml.document.main+xml"/></Types>',
        )
        z.writestr(
            "_rels/.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
            'officeDocument" Target="word/document.xml"/></Relationships>',
        )
        z.writestr(
            "word/document.xml",
            '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p><w:r><w:t>Quarterly staffing plan for Northwind</w:t></w:r></w:p></w:body></w:document>",
        )
    return str(path)


def _minimal_pptx(tmp_path, name="deck.pptx") -> str:
    """A real one-slide deck, authored with python-pptx.

    python-pptx arrives with ``markitdown[all]`` (the ``[extraction]`` extra),
    so where the markitdown route can run the author library is there too. A
    hand-written OOXML presentation needs a dozen interlocking parts to open
    at all, which is more fixture than test.
    """
    pptx = pytest.importorskip("pptx", reason="python-pptx ships with markitdown[all] (the extraction extra)")
    from pptx.util import Inches

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])  # title only
    slide.shapes.title.text = "Engagement type breakdown"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box.text_frame.text = "Fixed-fee engagements grew in the third quarter"
    path = tmp_path / name
    deck.save(str(path))
    return str(path)


@pytest.fixture
def stdlib_only(monkeypatch):
    """Pin extraction to the lightweight path.

    ``extract_text`` tries docling first, so on a ``-rich`` image the assertions
    below would describe whichever reader answered rather than the stdlib one
    they are about. Docling does not read EPUB today — this keeps that from
    being the thing the test silently depends on.
    """
    from src.ingest import text_extract

    monkeypatch.setattr(text_extract, "_try_docling", lambda path: None)


def _minimal_epub(tmp_path, name="book.epub", *, spine=True, chapters=None) -> str:
    """A real (if minimal) EPUB — a zip carrying the parts a reader needs.

    ``chapters`` is a list of ``(href, body_text)`` in the order the spine
    should present them; the FILE names are chosen by the caller so a test can
    make spine order disagree with archive/alphabetical order. With
    ``spine=False`` the container/OPF pair is omitted, which is the malformed
    shape the fallback path exists for.
    """
    chapters = chapters or [("b-first.xhtml", "Chapter one text"), ("a-second.xhtml", "Chapter two text")]
    path = tmp_path / name
    with zipfile.ZipFile(path, "w") as z:
        # Per spec the mimetype entry is first and stored, not deflated.
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        if spine:
            z.writestr(
                "META-INF/container.xml",
                '<?xml version="1.0"?><container version="1.0" '
                'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
                '<rootfile full-path="OEBPS/content.opf" '
                'media-type="application/oebps-package+xml"/></rootfiles></container>',
            )
            manifest = "".join(
                f'<item id="c{i}" href="{href}" media-type="application/xhtml+xml"/>'
                for i, (href, _) in enumerate(chapters)
            )
            itemrefs = "".join(f'<itemref idref="c{i}"/>' for i in range(len(chapters)))
            z.writestr(
                "OEBPS/content.opf",
                '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
                '<metadata><dc:title xmlns:dc="http://purl.org/dc/elements/1.1/">Northwind handbook'
                f"</dc:title></metadata><manifest>{manifest}</manifest><spine>{itemrefs}</spine></package>",
            )
        for href, body in chapters:
            z.writestr(
                f"OEBPS/{href}",
                '<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"><head>'
                "<title>Chapter title</title><style>p { color: red }</style></head>"
                f"<body><p>{body}</p></body></html>",
            )
    return str(path)


# ---------------------------------------------------------------------------
# Without the extra: the degradation has to say what it is
# ---------------------------------------------------------------------------


@pytest.mark.skipif(not _HAS_MARKITDOWN, reason="needs the [extraction] extra (default image, every default CI job)")
@pytest.mark.parametrize("ext", ["docx", "pptx"])
def test_office_documents_index_through_markitdown_on_the_default_image(tmp_path, ext):
    """The bug this pins: a ``.pptx`` uploaded to a collection was stored and
    then rejected with "needs the 'docling' extra" — on an image that already
    carried markitdown for the SharePoint crawl, which converts the very same
    file type fine. Office documents must index wherever the ``[extraction]``
    extra is, and the default image ships it."""
    if ext == "docx":
        path, expected = _minimal_docx(tmp_path), "Northwind"
    else:
        path, expected = _minimal_pptx(tmp_path), "Engagement type breakdown"

    res = extract_text(path, ext)

    assert expected in res.full_text, f"a valid .{ext} yielded no text on an image with markitdown installed"


@pytest.mark.parametrize("ext", ["docx", "pptx"])
def test_office_formats_name_both_extras_when_neither_parser_is_installed(tmp_path, monkeypatch, ext):
    """These formats are on the upload allowlist (``TIER1_EXTENSIONS``) and
    have no stdlib fallback, so on an image with neither parser a user
    uploads successfully and the file is rejected afterwards. A bare "no text
    extractor for '.docx'" reads as "your file is broken" and invites the one
    action that cannot help — uploading it again. Simulated, because the
    default install HAS markitdown: this pins the degraded deployment."""
    import sys

    from src.ingest import convert

    monkeypatch.setitem(sys.modules, "markitdown", None)
    monkeypatch.setattr(convert, "docling_capability", lambda: False)
    path = tmp_path / f"a.{ext}"
    with zipfile.ZipFile(path, "w") as package:  # a zip, so the archive guard is not what refuses it
        package.writestr("[Content_Types].xml", "<Types/>")

    with pytest.raises(UnsupportedDocument) as exc:
        extract_text(str(path), ext)

    msg = str(exc.value)
    assert "extraction" in msg and "docling" in msg, "the rejection must name BOTH extras that could read it"
    assert "built without" in msg, "…and say it is the deployment's shape, not the file's"
    assert "-rich" in msg and "DEPLOYMENT.md" in msg, "…and point an operator at the fix"


@pytest.mark.skipif(not _HAS_MARKITDOWN, reason="needs the [extraction] extra (default image, every default CI job)")
def test_corrupt_office_file_rejects_without_blaming_a_missing_extra(tmp_path):
    """A truncated archive is the file's problem. Naming an extra would send
    an operator to rebuild an image that changes nothing.

    The bytes carry NULs on purpose: markitdown sniffs content, and a
    "corrupt" archive that is really ASCII prose is read as prose — which is
    the right outcome for that file, and not the case under test."""
    path = tmp_path / "a.docx"
    path.write_bytes(b"PK\x03\x04\x00\x00truncated archive")

    with pytest.raises(UnsupportedDocument) as exc:
        extract_text(str(path), "docx")

    msg = str(exc.value)
    assert "built without" not in msg and "-rich" not in msg
    assert "docx" in msg


def test_allowlisted_formats_without_a_reader_are_exactly_the_ones_we_name():
    """Guard against drift between the upload allowlist and the extraction
    story: any tier-1 extension that neither the allowlist-independent
    fallbacks nor docling-only set covers would reject with the old generic
    message, i.e. an accepted upload with an unexplained failure."""
    from src.corpus_allowlist import TIER1_EXTENSIONS
    from src.ingest.runner import TABULAR_EXTS
    from src.ingest.text_extract import (
        _HTML_EXTS,
        _OFFICE_EXTS,
        _PLAIN_EXTS,
        _UNREADABLE_TIER1,
    )

    covered = (
        set(_PLAIN_EXTS)
        | set(_HTML_EXTS)
        | set(TABULAR_EXTS)
        | set(_OFFICE_EXTS)
        | set(_UNREADABLE_TIER1)
        | {"pdf", "eml", "epub"}
    )
    unexplained = set(TIER1_EXTENSIONS) - covered
    assert unexplained == set(), (
        f"tier-1 uploads with no reader and no named reason: {sorted(unexplained)} — either add a "
        "fallback, add them to _OFFICE_EXTS so the shared converter reads them and the rejection "
        "names the missing extras, or record them in _UNREADABLE_TIER1 (with the reason) if nothing "
        "can read them"
    )


def test_eml_is_parsed_by_the_stdlib_on_every_build(tmp_path):
    """``.eml`` is allowlisted and used to be rejected on every build — no
    reader existed, with or without extras. Email threads are ordinary
    corpus material, so this one is parsed with the stdlib rather than
    behind an extra."""
    path = tmp_path / "thread.eml"
    path.write_text(
        "From: partner@example.com\n"
        "To: team@example.com\n"
        "Subject: Northwind staffing\n"
        "Date: Mon, 4 Aug 2026 09:12:00 +0200\n"
        "\n"
        "Confirming two analysts for the September phase.\n",
        encoding="utf-8",
    )
    res = extract_text(str(path), "eml")
    assert "Northwind staffing" in res.full_text, "the subject must be searchable"
    assert "partner@example.com" in res.full_text, "so must the correspondents"
    assert "two analysts" in res.full_text, "and the body"


def test_epub_is_parsed_on_every_build(tmp_path):
    """``.epub`` is allowlisted and used to be rejected on every build — no
    reader existed, with or without extras. An EPUB is a zip of XHTML, so this
    one is parsed with the stdlib rather than behind an extra, the same call
    made for ``.eml``. Deliberately NOT pinned to the stdlib path: the claim is
    that an EPUB yields text on any image."""
    res = extract_text(_minimal_epub(tmp_path), "epub")

    assert "Chapter one text" in res.full_text
    assert "Chapter two text" in res.full_text


def test_epub_indexes_the_title_and_drops_stylesheets(tmp_path, stdlib_only):
    """A chapter's ``<title>`` is ordinary searchable text and stays. A
    ``<style>`` body is markup that would only pollute the index."""
    path = _minimal_epub(tmp_path, chapters=[("ch.xhtml", "Body prose here")])
    text = extract_text(path, "epub").full_text

    assert "Body prose here" in text
    assert "chapter title" in text.lower(), "the document title is searchable content"
    assert "color: red" not in text, "<style> content must not be indexed"


def test_epub_reading_order_follows_the_spine_not_the_archive(tmp_path, stdlib_only):
    """Chunking gives a reader neighbouring text as context, so chapter order
    is load-bearing. The spine is the only thing that states it — file names
    inside an EPUB are arbitrary, and here they sort the other way round."""
    path = _minimal_epub(
        tmp_path,
        chapters=[("z-opening.xhtml", "The opening chapter"), ("a-closing.xhtml", "The closing chapter")],
    )
    text = extract_text(path, "epub").full_text

    assert text.index("opening chapter") < text.index("closing chapter"), (
        "spine order was ignored — chapters came back in archive/alphabetical order"
    )


# Legal OPF spellings a real EPUB may use. Each is a way of writing the SAME
# manifest+spine; missing one does not fail loudly, it drops the spine and
# falls back to archive order — chapters in an arbitrary sequence, with
# nothing in the output saying so. Found by probing variants, not by reading
# the spec, which is why they are pinned here.
_OPF_VARIANTS = {
    "double-quoted": (
        '<package><manifest><item id="a" href="a-two.xhtml"/><item id="b" href="z-one.xhtml"/>'
        '</manifest><spine><itemref idref="a"/><itemref idref="b"/></spine></package>',
        '"',
    ),
    "single-quoted": (
        "<package><manifest><item id='a' href='a-two.xhtml'/><item id='b' href='z-one.xhtml'/>"
        "</manifest><spine><itemref idref='a'/><itemref idref='b'/></spine></package>",
        "'",
    ),
    "namespace-prefixed": (
        '<opf:package xmlns:opf="http://www.idpf.org/2007/opf"><opf:manifest>'
        '<opf:item id="a" href="a-two.xhtml"/><opf:item id="b" href="z-one.xhtml"/>'
        '</opf:manifest><opf:spine><opf:itemref idref="a"/><opf:itemref idref="b"/>'
        "</opf:spine></opf:package>",
        '"',
    ),
    "decoy-data-id-attribute": (
        '<package><manifest><item data-id="wrong" id="a" href="a-two.xhtml"/>'
        '<item id="b" href="z-one.xhtml"/></manifest>'
        '<spine><itemref idref="a"/><itemref idref="b"/></spine></package>',
        '"',
    ),
}


@pytest.mark.parametrize("variant", sorted(_OPF_VARIANTS))
def test_epub_spine_survives_legal_opf_spellings(tmp_path, stdlib_only, variant):
    """Archive order is ``z-one`` then ``a-two``; every spine here says the
    OPPOSITE. So ``a-two`` coming out first is the only evidence the spine was
    actually parsed — with one chapter, or with chapters in agreeing order,
    the fallback would look identical and the test would prove nothing."""
    opf_body, quote = _OPF_VARIANTS[variant]
    path = tmp_path / f"{variant}.epub"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr(
            "META-INF/container.xml",
            '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            f"<rootfiles><rootfile full-path={quote}OEBPS/content.opf{quote} "
            'media-type="application/oebps-package+xml"/></rootfiles></container>',
        )
        z.writestr("OEBPS/content.opf", opf_body)
        z.writestr("OEBPS/z-one.xhtml", "<html><body><p>ARCHIVE FIRST</p></body></html>")
        z.writestr("OEBPS/a-two.xhtml", "<html><body><p>SPINE FIRST</p></body></html>")

    text = extract_text(str(path), "epub").full_text

    assert text.index("SPINE FIRST") < text.index("ARCHIVE FIRST"), (
        f"the {variant} OPF was not parsed — extraction silently fell back to archive order"
    )


def test_epub_without_a_usable_spine_still_yields_text(tmp_path, stdlib_only):
    """EPUBs in the wild are malformed. A missing container/OPF pair must
    degrade to 'every XHTML member in archive order', not to a rejection —
    the text is right there."""
    res = extract_text(_minimal_epub(tmp_path, spine=False), "epub")

    assert "Chapter one text" in res.full_text
    assert "Chapter two text" in res.full_text


def test_epub_decompression_is_capped(tmp_path, monkeypatch, stdlib_only):
    """An EPUB is an untrusted zip: a few KB of deflate can inflate to
    gigabytes. Extraction stops at a byte ceiling instead of reading whatever
    the archive claims."""
    from src.ingest import text_extract

    monkeypatch.setattr(text_extract, "_EPUB_MAX_TEXT_BYTES", 2048)
    path = tmp_path / "bomb.epub"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
        z.writestr("OEBPS/huge.xhtml", "<html><body><p>" + ("A" * 8_000_000) + "</p></body></html>")

    res = text_extract.extract_text(str(path), "epub")

    assert len(res.full_text) <= 2048 * 2, f"cap ignored — got {len(res.full_text)} chars"


def test_broken_epub_rejects_without_blaming_a_missing_extra(tmp_path, stdlib_only):
    """A corrupt archive is the file's problem. Naming an extra would send an
    operator to rebuild an image that changes nothing."""
    path = tmp_path / "corrupt.epub"
    path.write_bytes(b"PK\x03\x04 truncated garbage")

    with pytest.raises(UnsupportedDocument) as exc:
        extract_text(str(path), "epub")

    assert "docling" not in str(exc.value)
    assert "epub" in str(exc.value).lower()


def test_unreadable_allowlisted_formats_still_reject_clearly(tmp_path):
    """A format nothing can read must not claim an extra would fix it.

    The set is EMPTY as of the epub reader landing, so the loop below is
    vacuous — which is why the membership assertion comes first. It catches
    the mistake that emptying this set invites: dropping an extension from
    the upload allowlist and leaving its name here, where it would describe
    a rejection path no upload can reach."""
    from src.corpus_allowlist import TIER1_EXTENSIONS
    from src.ingest.text_extract import _UNREADABLE_TIER1

    assert set(_UNREADABLE_TIER1) <= set(TIER1_EXTENSIONS), (
        "_UNREADABLE_TIER1 names a format the upload allowlist no longer accepts: "
        f"{sorted(set(_UNREADABLE_TIER1) - set(TIER1_EXTENSIONS))} — such an entry is dead "
        "weight describing a path nothing reaches; drop it"
    )

    for ext in sorted(_UNREADABLE_TIER1):
        path = tmp_path / f"a.{ext}"
        path.write_bytes(b"\x00binary")
        with pytest.raises(UnsupportedDocument) as exc:
            extract_text(str(path), ext)
        assert "docling" not in str(exc.value), (
            f"'.{ext}' cannot be read even with the extra — the message must not send an "
            "operator to rebuild the image for nothing"
        )


def test_capability_probe_does_not_import_the_heavy_dependency(monkeypatch):
    """``docling_capability`` is asked while composing an error message and
    while answering readiness — it must stay an import-spec probe, never an
    import, or a rejection path would pay torch's load cost."""
    import sys

    before = set(sys.modules)
    docling_capability()
    newly_imported = {m for m in set(sys.modules) - before if m.startswith("docling")}
    assert newly_imported == set(), f"capability probe imported {newly_imported}"


# ---------------------------------------------------------------------------
# With the extras: the capability is real
# ---------------------------------------------------------------------------


@pytest.mark.skipif(not _HAS_DOCLING, reason="needs the [docling] extra (rich image / rich-extras CI job)")
def test_docx_ingests_when_the_extra_is_present(tmp_path):
    """The whole point of the rich image: an office document yields text
    rather than a rejection."""
    res = extract_text(_minimal_docx(tmp_path), "docx")
    assert res.full_text.strip(), "docling returned no text for a valid docx"
    assert "Northwind" in res.full_text


def test_retrieval_mode_matches_the_installed_extra():
    """``retrieval_mode`` is the label a client reads to tell hybrid results
    from degraded ones, so it must track what is installed rather than a
    config value someone can set optimistically."""
    from src.ingest.embeddings import embedding_capability
    from src.ingest.retrieval import retrieval_mode

    expected = "hybrid" if embedding_capability() else "lexical_only"
    assert retrieval_mode() == expected
