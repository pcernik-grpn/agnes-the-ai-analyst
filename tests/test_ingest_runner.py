"""Tests for src.ingest.runner.ingest_file + tabular contract output."""

from __future__ import annotations

from pathlib import Path

import pytest


def _new_corpus(slug: str) -> str:
    from src.repositories import file_corpora_repo

    return file_corpora_repo().create(name=slug, slug=slug, description=None, created_by="u1")


def _add_file(corpus_id: str, filename: str, file_type: str, path: str) -> str:
    from src.repositories import corpus_files_repo

    return corpus_files_repo().add(
        corpus_id=corpus_id,
        filename=filename,
        sha256="sha_" + filename,
        file_type=file_type,
        size_bytes=Path(path).stat().st_size if Path(path).exists() else 0,
        storage_path=path,
    )


def test_ingest_csv_indexes_as_registered_table(e2e_env, tmp_path):
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_files_repo

    corpus_id = _new_corpus("ing-csv")
    csv = tmp_path / "data.csv"
    csv.write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
    file_id = _add_file(corpus_id, "sales data.csv", "csv", str(csv))

    assert ingest_file(file_id) == "indexed"
    row = corpus_files_repo().get(file_id)
    assert row["processing_status"] == "indexed"
    detail = row["processing_detail"]
    assert detail["kind"] == "tabular"
    table_id = detail["derived_table_id"]
    assert table_id

    # Contract output: parquet written + registered in table_registry.
    import os

    from src.repositories import table_registry_repo

    parquet = (
        Path(os.environ.get("DATA_DIR", "data"))
        / "extracts"
        / f"collection_{corpus_id}"
        / "data"
        / f"{table_id}.parquet"
    )
    assert parquet.exists()
    reg = table_registry_repo().get(table_id)
    assert reg is not None
    assert reg["query_mode"] == "local"


def test_ingest_csv_derived_table_is_owner_private(e2e_env, tmp_path):
    """#4: the uploader owns the derived table (registered_by) and can read it
    via RBAC, while another user cannot — access tracks the owning collection."""
    from src.ingest.runner import ingest_file
    from src.rbac import can_access_table, get_accessible_tables
    from src.repositories import corpus_files_repo, table_registry_repo, users_repo

    users_repo().create(id="u1", email="u1@test.com", name="Owner")
    users_repo().create(id="u2", email="u2@test.com", name="Other")

    corpus_id = _new_corpus("ing-priv")  # created_by="u1"
    csv = tmp_path / "priv.csv"
    csv.write_text("a,b\n1,2\n", encoding="utf-8")
    file_id = _add_file(corpus_id, "priv.csv", "csv", str(csv))
    assert ingest_file(file_id) == "indexed"

    table_id = corpus_files_repo().get(file_id)["processing_detail"]["derived_table_id"]
    # Provenance: the real uploader, not the literal "ingest".
    assert table_registry_repo().get(table_id)["registered_by"] == "u1"

    # RBAC: owner in, other out.
    assert can_access_table({"id": "u1"}, table_id) is True
    assert can_access_table({"id": "u2"}, table_id) is False
    assert table_id in get_accessible_tables({"id": "u1"})
    assert table_id not in get_accessible_tables({"id": "u2"})


def test_ingest_txt_creates_chunks(e2e_env, tmp_path):
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_chunks_repo, corpus_files_repo

    corpus_id = _new_corpus("ing-txt")
    doc = tmp_path / "notes.txt"
    doc.write_text("paragraph one.\n\nparagraph two has more text.", encoding="utf-8")
    file_id = _add_file(corpus_id, "notes.txt", "txt", str(doc))

    assert ingest_file(file_id) == "indexed"
    row = corpus_files_repo().get(file_id)
    assert row["processing_detail"]["kind"] == "document"
    chunks = corpus_chunks_repo().list_for_file(file_id)
    assert len(chunks) >= 1
    assert row["processing_detail"]["chunk_count"] == len(chunks)


def test_ingest_uses_preloaded_text_and_skips_the_disk_re_read(e2e_env, tmp_path):
    """A caller that already has the document's text in memory (the
    SharePoint crawl pipeline: `_prepare_document` converts, then
    `_Ingestor.ingest` writes it and calls `ingest_file`) must not pay a
    redundant read of the same content back off disk — a real, measured
    contributor to the crawl's parent-process memory pressure (a second
    full-size copy of the converted markdown, on top of every copy already
    held by convert/anonymize/encode/store). Proven here by pointing
    `storage_path` at a file that does not exist at all: if `ingest_file`
    ever fell back to reading it, this fails loudly instead of indexing the
    preloaded text.
    """
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_chunks_repo, corpus_files_repo

    corpus_id = _new_corpus("ing-preloaded")
    missing_path = str(tmp_path / "does-not-exist.md")
    file_id = _add_file(corpus_id, "doc.md", "md", missing_path)

    status = ingest_file(file_id, preloaded_text="already-converted markdown body")
    assert status == "indexed"
    row = corpus_files_repo().get(file_id)
    assert row["processing_status"] == "indexed"
    chunks = corpus_chunks_repo().list_for_file(file_id)
    assert len(chunks) == 1
    assert chunks[0]["text"] == "already-converted markdown body"


def test_ingest_image_stays_pending_for_vision_slice(e2e_env, tmp_path, monkeypatch):
    from src.ingest import vision

    # Force vision off for determinism (a dev with ANTHROPIC_API_KEY set would
    # otherwise make a real API call here).
    monkeypatch.setattr(vision, "extract_image_text", lambda path, *, ext: None)
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_files_repo

    corpus_id = _new_corpus("ing-img")
    img = tmp_path / "pic.png"
    img.write_bytes(b"\x89PNG\r\n\x1a\n fake")
    file_id = _add_file(corpus_id, "pic.png", "png", str(img))

    assert ingest_file(file_id) == "pending"
    assert corpus_files_repo().get(file_id)["processing_detail"]["tier"] == 2


def test_ingest_unextractable_document_rejected(e2e_env, tmp_path):
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_files_repo

    corpus_id = _new_corpus("ing-rej")
    # A truncated OOXML package: no reader on any image can open it. NULs on
    # purpose — markitdown sniffs content, and ASCII bytes in a ".docx" would
    # be read as the prose they are rather than rejected.
    doc = tmp_path / "report.docx"
    doc.write_bytes(b"PK\x03\x04\x00\x00 not really a docx")
    file_id = _add_file(corpus_id, "report.docx", "docx", str(doc))

    assert ingest_file(file_id) == "rejected"
    assert "reason" in corpus_files_repo().get(file_id)["processing_detail"]


def test_ingest_idempotent_rechunk(e2e_env, tmp_path):
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_chunks_repo

    corpus_id = _new_corpus("ing-idem")
    doc = tmp_path / "a.md"
    doc.write_text("# H\n\nsome content here", encoding="utf-8")
    file_id = _add_file(corpus_id, "a.md", "md", str(doc))

    ingest_file(file_id)
    first = len(corpus_chunks_repo().list_for_file(file_id))
    ingest_file(file_id)  # re-ingest must not duplicate
    second = len(corpus_chunks_repo().list_for_file(file_id))
    assert first == second


def test_tabular_same_base_name_distinct_tables(e2e_env, tmp_path):
    """Two files whose names sanitize to the same base must NOT collide on the
    derived DuckDB table (regression: silent overwrite)."""
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_files_repo

    corpus_id = _new_corpus("ing-collide")
    d1 = tmp_path / "one"
    d2 = tmp_path / "two"
    d1.mkdir()
    d2.mkdir()
    (d1 / "data.csv").write_text("region,revenue\nEU,100\n", encoding="utf-8")
    (d2 / "data.csv").write_text("region,revenue\nUS,250\n", encoding="utf-8")
    fid1 = _add_file(corpus_id, "data.csv", "csv", str(d1 / "data.csv"))
    fid2 = _add_file(corpus_id, "data.csv", "csv", str(d2 / "data.csv"))

    assert ingest_file(fid1) == "indexed"
    assert ingest_file(fid2) == "indexed"
    t1 = corpus_files_repo().get(fid1)["processing_detail"]["derived_table_id"]
    t2 = corpus_files_repo().get(fid2)["processing_detail"]["derived_table_id"]
    assert t1 != t2, f"same-base files collided on table_id: {t1}"


def test_tabular_reingest_is_idempotent(e2e_env, tmp_path):
    """Re-ingesting the same tabular file must not raise — table_registry.register()
    upserts on id (ON CONFLICT DO UPDATE), and parquet/_meta/view are overwrite-safe."""
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_files_repo, table_registry_repo

    corpus_id = _new_corpus("ing-tab-idem")
    csv = tmp_path / "data.csv"
    csv.write_text("region,revenue\nEU,100\n", encoding="utf-8")
    fid = _add_file(corpus_id, "data.csv", "csv", str(csv))

    assert ingest_file(fid) == "indexed"
    t1 = corpus_files_repo().get(fid)["processing_detail"]["derived_table_id"]
    # Second pass (e.g. retry) must succeed, not raise a unique-constraint error.
    assert ingest_file(fid) == "indexed"
    t2 = corpus_files_repo().get(fid)["processing_detail"]["derived_table_id"]
    assert t1 == t2
    assert table_registry_repo().get(t1) is not None


def test_empty_tabular_is_needs_review_not_indexed(e2e_env, tmp_path):
    """Header-only CSV → 0 rows: must NOT register a table nor claim indexed."""
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_files_repo, table_registry_repo

    corpus_id = _new_corpus("ing-empty-csv")
    csv = tmp_path / "empty.csv"
    csv.write_text("a,b\n", encoding="utf-8")
    file_id = _add_file(corpus_id, "empty.csv", "csv", str(csv))

    assert ingest_file(file_id) == "needs_review"
    row = corpus_files_repo().get(file_id)
    assert row["processing_status"] == "needs_review"
    assert "empty" in row["processing_detail"]["reason"]
    # No derived table may leak into the registry.
    fid_suffix = file_id.replace("cf_", "")[:8]
    leaked = [r for r in table_registry_repo().list_by_source("collection") if r.get("id", "").endswith(fid_suffix)]
    assert leaked == []


def test_zero_chunk_document_is_needs_review(e2e_env, tmp_path, monkeypatch):
    """Extractor succeeds but yields no text → needs_review, not indexed."""
    import src.ingest.runner as runner_mod
    from src.repositories import corpus_files_repo

    corpus_id = _new_corpus("ing-zero-chunks")
    doc = tmp_path / "blank.txt"
    doc.write_text("", encoding="utf-8")
    file_id = _add_file(corpus_id, "blank.txt", "txt", str(doc))

    monkeypatch.setattr(runner_mod, "extract_text", lambda p, t: "")
    assert runner_mod.ingest_file(file_id) == "needs_review"
    row = corpus_files_repo().get(file_id)
    assert row["processing_status"] == "needs_review"
    assert row["processing_detail"]["reason"] == "extraction produced no text chunks"


def test_ingest_preloaded_text_with_a_nul_byte_is_indexed_and_searchable(e2e_env, tmp_path):
    """A NUL byte mid-word in converted text (a SharePoint crawl's own path:
    `preloaded_text`) must not reject the document — PostgreSQL `text`
    columns refuse to store `0x00` outright, which surfaced live as 261
    rejected documents (an Oracle table export, several ordinary
    SharePoint files) before `src.ingest.chunking._sanitize_control_chars`
    started stripping it at the ingest boundary. Proven end to end: the
    document indexes (not rejects) and a term either side of the stripped
    byte finds it back through `src.ingest.retrieval.search`.
    """
    from src.ingest.retrieval import search
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_chunks_repo, corpus_files_repo

    corpus_id = _new_corpus("ing-nul")
    missing_path = str(tmp_path / "does-not-exist.md")
    file_id = _add_file(corpus_id, "ap_suppliers.md", "md", missing_path)

    status = ingest_file(file_id, preloaded_text="in.c_keboola_ex_db_or\x00acle_ap_suppliers table export")
    assert status == "indexed"
    row = corpus_files_repo().get(file_id)
    assert row["processing_status"] == "indexed"

    chunks = corpus_chunks_repo().list_for_file(file_id)
    assert len(chunks) == 1
    assert "\x00" not in chunks[0]["text"]
    assert chunks[0]["text"] == "in.c_keboola_ex_db_oracle_ap_suppliers table export"

    results = search([corpus_id], "oracle_ap_suppliers")
    assert any(r["file_id"] == file_id for r in results)


def test_ingest_file_routes_zip_to_bundle(e2e_env, tmp_path, monkeypatch):
    """A zip row delegates to ingest_bundle (K1) instead of the prose path."""
    import src.ingest.bundle as bundle_mod
    from src.ingest.runner import ingest_file

    corpus_id = _new_corpus("ing-zip-route")
    zip_path = tmp_path / "dump.zip"
    zip_path.write_bytes(b"PK\x03\x04placeholder")
    file_id = _add_file(corpus_id, "dump.zip", "zip", str(zip_path))

    calls: dict = {}

    def fake_bundle(cid, fid, path, **kw):
        calls["args"] = (cid, fid, path)
        return "indexed"

    monkeypatch.setattr(bundle_mod, "ingest_bundle", fake_bundle)
    assert ingest_file(file_id) == "indexed"
    assert calls["args"] == (corpus_id, file_id, str(zip_path))


def test_chunks_land_in_the_collection_the_file_is_in_when_they_are_written(e2e_env, tmp_path, monkeypatch):
    """A move that happens WHILE a file is being ingested must not be undone
    by the ingest finishing.

    ``ingest_file`` reads the file's collection once at the top and then does
    the slow part (extraction, OCR, embedding). A move landing in that window
    re-homes the chunks that exist at the time and cannot touch rows the
    ingest has yet to write — so the ingest used to write its chunks under
    the collection the file has already left, restoring the very leak the
    move exists to close. The collection is re-read at write time instead.
    """
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_chunks_repo, corpus_files_repo

    src_id = _new_corpus("ing-move-src")
    dst_id = _new_corpus("ing-move-dst")
    doc = tmp_path / "moving.txt"
    doc.write_text("the body of a file that moves mid-ingest", encoding="utf-8")
    file_id = _add_file(src_id, "moving.txt", "txt", str(doc))

    real_extract = __import__("src.ingest.runner", fromlist=["extract_text"]).extract_text

    def _extract_then_move(path, file_type):
        # The move lands after the collection was read, before chunks exist.
        corpus_files_repo().move_to_corpus(file_id, dst_id)
        corpus_chunks_repo().reassign_file_corpus(file_id, dst_id)
        return real_extract(path, file_type)

    monkeypatch.setattr("src.ingest.runner.extract_text", _extract_then_move)
    assert ingest_file(file_id) == "indexed"

    assert corpus_files_repo().get(file_id)["corpus_id"] == dst_id
    written = {ch["corpus_id"] for ch in corpus_chunks_repo().list_for_file(file_id)}
    assert written == {dst_id}
    assert corpus_chunks_repo().list_for_corpus(src_id) == []


def _write_pptx_with_one_picture(path: Path) -> Path:
    """A real ``.pptx`` (python-pptx — an existing markitdown dependency,
    never a new one) with a single embedded picture on its one slide."""
    import io
    import struct
    import zlib

    from pptx import Presentation
    from pptx.util import Inches

    def _chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    png = (
        b"\x89PNG\r\n\x1a\n"
        + _chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + _chunk(b"IDAT", zlib.compress(b"\x00\x00\x00\x00"))
        + _chunk(b"IEND", b"")
    )
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(1))
    presentation.save(str(path))
    return path


def test_ingest_pptx_with_embedded_picture_discloses_it_in_the_chunk_and_the_detail(e2e_env, tmp_path):
    """The live finding this pins (2026-09-09): a reader asked about content
    that lives in a slide's picture and got an answer assembled from prose
    instead, with no indication the picture existed at all. The indexed
    chunk must carry an honest, located disclosure instead of a bare
    ``PictureN.jpg`` placeholder (a name that collides across documents —
    see ``src/ingest/convert.py``'s module docstring), and the file's own
    ``processing_detail`` must carry the count so an admin/agent can act on
    it without re-scanning every chunk.
    """
    pytest.importorskip("markitdown", reason="extraction extra not installed")
    from src.ingest.runner import ingest_file
    from src.repositories import corpus_chunks_repo, corpus_files_repo

    corpus_id = _new_corpus("ing-pptx-image")
    deck = _write_pptx_with_one_picture(tmp_path / "architecture.pptx")
    file_id = _add_file(corpus_id, "architecture.pptx", "pptx", str(deck))

    assert ingest_file(file_id) == "indexed"

    row = corpus_files_repo().get(file_id)
    assert row["processing_detail"]["image_count"] == 1

    chunks = corpus_chunks_repo().list_for_file(file_id)
    full_text = "\n".join(c["text"] for c in chunks)
    assert ".jpg" not in full_text
    assert "[image 1 of 1 in this document — not indexed, slide 1]" in full_text
